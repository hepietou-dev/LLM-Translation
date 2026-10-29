"""DOCX builder: block IR -> formatted .docx (3 typography presets)."""
from __future__ import annotations

import re
from pathlib import Path

import re
import re as _re

from . import footnotes

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

C_BODY = RGBColor(0x22, 0x22, 0x22)
C_GRAY = RGBColor(0x66, 0x66, 0x66)
C_LINK = RGBColor(0x25, 0x63, 0xEB)

PRESETS = {
    # 屏幕阅读：无首行缩进、段间留白，适合长时间屏幕阅读
    "screen": dict(label="屏幕阅读", zh="微软雅黑", en="Segoe UI", size=11, spacing=1.6, after=9,
                   align="both", indent=0, h1=21, h2=15, h3=12.5, hcolor="1F3864",
                   rule=False, chapter_break=False, margin=(2.2, 2.4), header="title",
                   hbefore=(30, 22, 16), hafter=(14, 9, 7)),
    # 印刷版式：宋体正文、首行缩进两字、两端对齐、章另起页，接近学术出版物
    "print": dict(label="印刷版式", zh="宋体", en="Cambria", size=10.5, spacing=1.45, after=0,
                  align="both", indent=2, h1=20, h2=14, h3=12, hcolor="000000",
                  rule=False, chapter_break=True, margin=(2.5, 2.5), header="title",
                  hbefore=(0, 20, 14), hafter=(18, 10, 6)),
    # 大字版：行距与字号最大，照顾老年读者
    "large": dict(label="大字版", zh="微软雅黑", en="Segoe UI", size=13, spacing=1.85, after=11,
                  align="both", indent=0, h1=24, h2=17, h3=14, hcolor="1F3864",
                  rule=False, chapter_break=True, margin=(2.2, 2.6), header="title",
                  hbefore=(32, 24, 18), hafter=(16, 11, 8)),
}
PRESET_LABELS = {k: v["label"] for k, v in PRESETS.items()}
LABEL_TO_PRESET = {v: k for k, v in PRESET_LABELS.items()}


_BAD = _re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

NOTE_HEAD = _re.compile(
    r"^(?:NOTES?|ENDNOTES?|注\s*释?|尾注|脚注|编者注)\s*[:：]?\s*$", _re.I)
REF_HEAD = _re.compile(
    r"^(?:REFERENCES?|BIBLIOGRAPHY|WORKS CITED|FURTHER READING|参考文献|参考书目)"
    r"\s*[:：]?\s*$", _re.I)

MARK = _re.compile(r"\u27e6(\d{1,3})\u27e7")
EMPH = _re.compile(r"\*([^*\n]+)\*")
REF = _re.compile(r"(?i)(图|表|Figure|Fig\.|Table|Map)\s*([0-9]+(?:\.[0-9]+)*)")
NOTE_NO = _re.compile(
    r"^\s*(?:\u27e6\s*(\d{1,3})\s*\u27e7|\[\s*(\d{1,3})\s*\]"
    r"|(\d{1,3})\s*[.、:：)）]\s*|\s*(\d{1,3})\s+)")
SRC_PAGE = _re.compile(r"[（(]?\s*原书\s*p\.\s*[\wivxlcdm-]+\s*[）)]?\s*$", _re.I)


def _ref_anchor(m, refs):
    """Return (key, printed label) for a figure/table reference, or None."""
    if not refs:
        return None
    word, num = m.group(1), m.group(2)
    kind = "tab" if word.lower() in ("表", "table") else "fig"
    key = kind + ":" + num
    anchor = refs.get(key)
    if not anchor:
        return None
    return anchor, ("表 " if kind == "tab" else "图 ") + num


def _add_runs(p, text, bold=False, italic=False, refs=None, size=None,
              en="Segoe UI", zh="微软雅黑"):
    """Plain run helper: *italic* spans and figure/table cross-references."""
    pos = 0
    for m in EMPH.finditer(text):
        if m.start() > pos:
            _add_ref_runs(p, text[pos:m.start()], refs, bold, italic, size, en, zh)
        r = p.add_run(m.group(1))
        r.bold, r.italic = bold, True
        pos = m.end()
    if pos < len(text):
        _add_ref_runs(p, text[pos:], refs, bold, italic, size, en, zh)


def _add_ref_runs(p, text, refs, bold, italic, size, en, zh):
    if not text:
        return
    pos = 0
    for m in REF.finditer(text):
        hit = _ref_anchor(m, refs)
        if not hit:
            continue
        anchor, label = hit
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()])
            r.bold, r.italic = bold, italic
        _link(p, anchor, label, size or 10.5, color=C_BODY, underline=False,
              bold=bold, italic=italic, en=en, zh=zh)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:])
        r.bold, r.italic = bold, italic


def _add_marked(p, text, bold=False, italic=False, on_note=None, refs=None,
                size=None, en="Segoe UI", zh="微软雅黑"):
    """Add text with ⟦n⟧ note calls as superscript runs (or live footnotes)."""
    text = re.sub(r"(\u27e6\d{1,3}\u27e7)[ \t]+(?=[\u4e00-\u9fff])", r"\1", text)
    pos = 0
    for m in MARK.finditer(text):
        if m.start() > pos:
            _add_runs(p, text[pos:m.start()], bold, italic, refs=refs, size=size, en=en, zh=zh)
        if not (on_note and on_note(m.group(1))):
            r = p.add_run(m.group(1))
            r.font.superscript = True
            r.bold, r.italic = bold, italic
        pos = m.end()
    if pos < len(text):
        _add_runs(p, text[pos:], bold, italic, refs=refs, size=size, en=en, zh=zh)


_BIB_YEAR = re.compile(
    r"\.\s+((?:1[3-9]|20)\d\d(?:[a-z]|(?:[–-]\d{1,4}))?"
    r"(?:,\s*\d{4}(?:[a-z]|(?:[–-]\d{1,4}))?)*)\.\s")
_BIB_END = re.compile(
    r"\.\s+(?:Ed\.|Edited by|Translated by|In\s|With\s)"
    r"|\.\s+[A-Z][^.:]{0,40}:"
    r"|\.\s+n\.p\."
    r"|,\s+\d+(?:st|nd|rd|th)\s+(?:edn|ed\.|edition|series)"
    r"|\.\s+\d+(?:st|nd|rd|th)\s+(?:edn|ed\.|edition)"
    r"|,\s+(?:\d+\s+vols?\.?|vols?\.?\s*\d+|Corrected and Expanded Edition|Expanded Edition)"
    r"|(?<=[”\"])\s+[A-Z][^.:]{0,40}:")
_BIB_IN = re.compile(
    r"^\s*,?\s*in\s+(.+?)(?=,\s*(?:\d|[ivxlcdm]+\s*[–-]|[Vv]ols?\.|[Vv]olume|Books?|Part"
    r"|Series\b)|\.\s+(?:Ed\.|Edited by|Translated by))")
_BIB_IN_PAGES = re.compile(r"^\s*,?\s*in\s+(.+?)(?=\.\s+\d)")
_BIB_IN_PLACE = re.compile(
    r"^\s*,?\s*in\s+(.+?)(?=\.\s+(?:Ed\.|Edited by|Translated by)|\.\s+[A-Z][^.:]{0,40}:)")
_BIB_JOURNAL = re.compile(r"^\s*[.,]?\s+([A-Z][^.]{1,90}?)\s+(?=\d)")
_BIB_EDIT_VOL = re.compile(r"^\s*,?\s+([A-Z][^.]{2,90}?)\.\s+\d+\s*[–-]\s*\d")
_BIB_ONLINE = re.compile(r"^\s*[.,]?\s+([A-Z][^.]{2,90}?)\.\s+(?=https?://)")
_BIB_SERIES = re.compile(r"^\s*[.,]?\s+([A-Z][^.]{2,90}?)\s+Ser\.\s*\d")
_BIB_JOURNAL_LATE = re.compile(r"\.\s+([A-Z][^.]{1,90}?)\s+\d+(?:\(\d+\))?:")
_BIB_TRIM = re.compile(r",\s+\d+(?:st|nd|rd|th)\s+series$|,\s*\d+\s+vols?\.?$")


def _bib_span(z: str):
    """Locate the book/journal title inside one reference entry."""
    m = _BIB_YEAR.search(z)
    if m:
        i = m.end()
    else:
        e = re.search(r"\b(?:ed|eds|trans)\.\s+", z)
        if not e:
            return None
        i = e.end()
    while True:
        r = re.match(r"Reprinted(?:\s+in)?\s+\d{4}(?:[–-]\d{2,4})?\.\s*", z[i:])
        if not r:
            break
        i += r.end()
    if z[i:i + 1] in "“\"":
        close = -1
        for q in re.finditer(r"[”\"]", z[i + 1:]):
            pos = i + 1 + q.start()
            if re.match(r"\s*[.,]?\s*(?:in\s|[A-Z]|\d)", z[pos + 1:]):
                close = pos
                break
        if close < 0:
            return None
        tail = z[close + 1:]
        for pat in (_BIB_IN, _BIB_IN_PAGES, _BIB_IN_PLACE, _BIB_JOURNAL,
                    _BIB_EDIT_VOL, _BIB_ONLINE, _BIB_SERIES):
            mm = pat.match(tail)
            if mm:
                s = close + 1 + mm.start(1)
                e = s + len(mm.group(1))
                mt = _BIB_TRIM.search(z[s:e])
                if mt:
                    e = s + mt.start()
                return s, e
        mm = _BIB_JOURNAL_LATE.search(tail)
        if mm:
            s = close + 1 + mm.start(1)
            return s, s + len(mm.group(1))
        return None
    mm = _BIB_END.search(z, i)
    if not mm:
        return None
    return i, mm.start()


def _bib_marked(z: str) -> str:
    """Wrap a reference entry's book/journal title in *…* for italics."""
    if not z or re.search(r"[\u4e00-\u9fff]", z):
        return z
    sp = _bib_span(z)
    if not sp:
        return z
    s, e = sp
    e = s + len(z[s:e].rstrip(" ,"))
    return z[:s] + "*" + z[s:e] + "*" + z[e:]


def toc_heads(blocks, meta=None):
    """Headings that belong in the printed contents.

    Chapter-end NOTES/REFERENCES headings repeat once per chapter; listing every
    one turns the TOC into noise. The same selection feeds the PDF page-number
    pass, so both sides stay aligned.
    """
    meta = meta or {}
    authors = [a for a in (meta.get("author_en"), meta.get("author"))
               if a and len(a) > 3]
    ref_heads = [b for b in blocks if b["type"] in ("h1", "h2")
                 and (REF_HEAD.match((b.get("text") or "").strip())
                      or REF_HEAD.match((b.get("zh") or "").strip()))]
    solo_ref = len(ref_heads) == 1
    out, seen = [], set()
    seen_h1 = False
    for b in blocks:
        if b["type"] not in ("h1", "h2"):
            continue
        orig = b["type"]
        t = (b.get("text") or "").strip()
        z = (b.get("zh") or "").strip()
        if NOTE_HEAD.match(t) or NOTE_HEAD.match(z):
            continue
        if (REF_HEAD.match(t) or REF_HEAD.match(z)) and not solo_ref:
            continue
        if orig == "h2" and b.get("endnote"):
            continue
        if z in ("目录", "目　录") or t.upper() in ("CONTENTS", "TABLE OF CONTENTS"):
            continue
        if any(a.lower() in t.lower() for a in authors):
            continue
        # front-matter sections (about the editors, contributors, preface) are
        # top-level even though the source sets them below chapter level
        if orig == "h2" and not seen_h1:
            b = dict(b, type="h1")
        if orig == "h1":
            seen_h1 = True
        if b["type"] == "h1":
            key = _re.sub(r"\s+", "", z or t)
            if key in seen:
                continue
            seen.add(key)
        out.append(b)
    return out


def _clean(s: str) -> str:
    return _BAD.sub(" ", s or "").replace("\ufffd", "").replace("\u00a0", " ")


def _fonts(rPr, en, zh):
    el = rPr.find(qn("w:rFonts"))
    if el is None:
        el = OxmlElement("w:rFonts")
        rStyle = rPr.find(qn("w:rStyle"))
        rStyle.addnext(el) if rStyle is not None else rPr.insert(0, el)
    for k, v in (("w:ascii", en), ("w:hAnsi", en), ("w:eastAsia", zh), ("w:cs", en)):
        el.set(qn(k), v)
    # theme fonts win over explicit names in LibreOffice
    for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        el.attrib.pop(qn(k), None)


def _style_font(style, en, zh, size=None, bold=None, color=None):
    _fonts(style.element.get_or_add_rPr(), en, zh)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = color


def _border(p, color="D9D9D9", sz=6, side="bottom"):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:" + side)
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz))
    e.set(qn("w:space"), "4"); e.set(qn("w:color"), color)
    pBdr.append(e)
    pPr.insert(0, pBdr)


def _bookmark(p, name, bid):
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bid)); s.set(qn("w:name"), name)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid))
    p._p.insert(0, s); p._p.append(e)


def _link(p, anchor, text, size, color="2563EB", underline=True,
          bold=False, italic=False, en="Segoe UI", zh="微软雅黑"):
    if not isinstance(color, str):
        color = str(color)                 # RGBColor -> "RRGGBB"
    hl = OxmlElement("w:hyperlink"); hl.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    _fonts(rPr, en, zh)
    c = OxmlElement("w:color"); c.set(qn("w:val"), color); rPr.append(c)
    for tag in ("w:sz", "w:szCs"):
        e = OxmlElement(tag); e.set(qn("w:val"), str(int(size * 2))); rPr.append(e)
    if bold:
        rPr.append(OxmlElement("w:b"))
    if italic:
        rPr.append(OxmlElement("w:i"))
    if underline:
        u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    r.append(t); hl.append(r); p._p.append(hl)


def _page_field(p):
    def run():
        return p.add_run()
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
    r = run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = " PAGE "
    r._r.append(it)
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)
    r = run(); r.text = "1"
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)
    for x in p.runs:
        x.font.size = Pt(9); x.font.color.rgb = C_GRAY


def _style_ref_field(p, style: str, fallback: str, size: float = 9):
    """STYLEREF field for the running chapter title; fallback keeps previews clean."""
    def run():
        return p.add_run()

    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "begin"); r._r.append(fc)
    r = run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve")
    it.text = ' STYLEREF "%s" \\* MERGEFORMAT ' % style
    r._r.append(it)
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "separate"); r._r.append(fc)
    r = run(); r.text = fallback
    r = run(); fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), "end"); r._r.append(fc)
    for x in p.runs:
        x.font.size = Pt(size); x.font.color.rgb = C_GRAY


def _sect(sec, margin):
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(margin[0])
    sec.left_margin = sec.right_margin = Cm(margin[1])
    sec.header_distance = sec.footer_distance = Cm(1.2)


def _first_line_chars(p, chars=200):
    ind = p._p.get_or_add_pPr().find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind"); p._p.get_or_add_pPr().append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), "420")


COPY_MARK = re.compile(r"(版权所有|保留所有权利|All rights reserved|©|ISSN|ISBN|DOI[:：]|"
                       r"年度评论出版社|Annu\. Rev\.|Copyright ©)", re.I)


def _is_legal_line(text: str) -> bool:
    """Journal masthead / copyright lines are not content."""
    t = (text or "").strip()
    return len(t) < 200 and bool(COPY_MARK.search(t))



def _page_tag_run(par, page, page_end=None):
    """Small grey tag showing the printed page(s) this paragraph covers."""
    if not page:
        return
    label = str(page)
    if page_end and str(page_end) != label:
        label = "%s–%s" % (page, page_end)
    r = par.add_run("〔p.%s〕 " % label)
    r.font.size = Pt(8)
    r.font.color.rgb = C_GRAY


def _note_number(raw: str) -> tuple[str, str]:
    m = NOTE_NO.match(raw or "")
    if not m:
        return "", (raw or "")
    no = next((g for g in m.groups() if g), "")
    return no, (raw or "")[m.end():].lstrip()


def _clean_note_text(text: str) -> str:
    text = re.sub(r"^【测试】\s*", "", text or "")
    return MARK.sub("", text).strip()


def _note_parts(text: str, zh: str) -> list[tuple[str, str]]:
    """Split a notes block into (number, text) pieces.

    The translation keeps the source ⟦n⟧ markers, so a block that merged
    several page-bottom notes is split on those markers first. Plain "n. text"
    blocks fall back to the numbered-list parser.
    """
    src = (text or "").strip()
    z = re.sub(r"^【测试】\s*", "", (zh or "").strip())
    raw = z or src
    marks = list(MARK.finditer(raw))
    if marks and (marks[0].start() <= 2 or len(marks) > 1):
        out = []
        for i, m in enumerate(marks):
            end = marks[i + 1].start() if i + 1 < len(marks) else len(raw)
            body = raw[m.end():end].strip()
            if body:
                out.append((m.group(1), body))
        if out:
            return out
    no, body = _note_number(raw)
    if not body and src and src != raw:
        no2, body = _note_number(src)
        no = no or no2
    out = []
    for piece in re.split(r"\s(?=\d{1,3}\s)", body):
        piece = piece.strip()
        if not piece:
            continue
        mm = re.match(r"^(\d{1,3})\s+(.*)$", piece, re.S)
        if mm and len(mm.group(2)) > 12:
            out.append((mm.group(1), mm.group(2).strip()))
        else:
            out.append((no, piece))
    return out


def _assign_figure_numbers(blocks):
    """Number figures/tables and build the cross-reference map.

    Explicit source numbers ("Figure 3.2", "图 3.2") are preserved. Captions
    without a number get the next free sequential number, so no figure is left
    as "图 原书 p.X" in a proof.
    """
    explicit = {"figure": set(), "table": set()}
    parsed = {}
    head = _re.compile(
        r"(?i)^(?:Figure|Fig\.|Map|Plate|Chart|Table|图|地图|表|插图|表格)\s*"
        r"([0-9]+(?:\.[0-9]+)*)\s*[.．:：、—-]?\s*(.*)$")
    for b in blocks:
        if b.get("type") not in ("figure", "table"):
            continue
        kind = "figure" if b["type"] == "figure" else "table"
        cap = _clean(b.get("caption") or "").strip()
        cap = _re.sub(r"^【测试】\s*", "", cap)
        m = head.match(cap)
        if m:
            explicit[kind].add(m.group(1))
            parsed[id(b)] = (kind, m.group(1), m.group(2).strip())
        else:
            parsed[id(b)] = (kind, "", cap)

    used = {k: set(v) for k, v in explicit.items()}
    seq = {}
    for k, nums in explicit.items():
        seq[k] = max([int(float(n)) for n in nums if n.replace(".", "", 1).isdigit()] or [0])
    refs, dups, n = {}, [], 0
    for b in blocks:
        if b.get("type") not in ("figure", "table"):
            continue
        kind, num, rest = parsed[id(b)]
        cn = "图" if kind == "figure" else "表"
        if not num:
            rest = SRC_PAGE.sub("", rest)
            rest = _re.sub(r"^(?:图|表)\s*[　:：.]?\s*", "", rest).strip()
            rest = _re.sub(r"^[（(]?\s*原书\s*p\.\s*[\wivxlcdm-]+\s*[）)]?\s*", "",
                           rest, flags=_re.I).strip()
            seq[kind] += 1
            while str(seq[kind]) in used[kind]:
                seq[kind] += 1
            num = str(seq[kind])
            used[kind].add(num)
        key = ("fig" if kind == "figure" else "tab") + ":" + num
        b["caption"] = cn + " " + num + (("　" + rest) if rest else "")
        b["refkey"] = key
        b["anchor"] = "_Fig%04d" % (2000 + n)
        b["bid"] = 2000 + n
        n += 1
        if key in refs:
            dups.append("%s %s" % (cn, num))
        refs.setdefault(key, b["anchor"])
    return refs, dups


def _collect_notes(blocks):
    """Index footnote blocks before rendering, so markers can bind by number."""
    endnotes = False
    endnote_ids = set()
    endnote_nos = set()
    recs = []
    for b in blocks:
        t = b.get("type")
        if t in ("h1", "h2"):
            head = (b.get("zh") or b.get("text") or "").strip()
            endnotes = bool(b.get("endnote")) or bool(NOTE_HEAD.match(head))
        if t != "note":
            continue
        if endnotes or b.get("endnote"):
            endnote_ids.add(id(b))
            for no, _text in _note_parts(b.get("text", ""), b.get("zh") or ""):
                if no:
                    endnote_nos.add(str(no))
            continue
        for no, text in _note_parts(b.get("text", ""), b.get("zh") or ""):
            text = _clean_note_text(text)
            if text:
                recs.append({"no": str(no or ""), "text": text,
                             "page": str(b.get("page") or ""), "used": False})
    return recs, endnote_ids, endnote_nos


def _page_distance(a: str, b: str) -> int:
    try:
        return abs(int(a) - int(b))
    except (TypeError, ValueError):
        return 999


def _note_lookup(recs, no, page):
    cands = [r for r in recs if not r["used"] and str(r["no"]) == str(no)]
    if not cands:
        return None, None
    same = [r for r in cands if str(r["page"]) == str(page)]
    if same:
        return same[0], "exact"
    near = [r for r in cands if _page_distance(r["page"], page) <= 1]
    if near:
        return min(near, key=lambda r: _page_distance(r["page"], page)), "near"
    # A unique note number may still be the right one when the source has no
    # numeric pagination (EPUB/roman front matter). With real page numbers a
    # distant match is more likely a different chapter's footnote.
    if not str(page).isdigit():
        loose = [r for r in cands if not str(r["page"]).isdigit()]
        if len(loose) == 1:
            return loose[0], "global"
    return None, None


def build_doc(book: dict, preset: str, out: Path, meta: dict, pages: dict | None = None,
              stats: dict | None = None) -> Path:
    cfg = PRESETS[preset]
    page_tags = bool(meta.get("page_tags", False))
    doc = Document()
    normal = doc.styles["Normal"]
    _style_font(normal, cfg["en"], cfg["zh"], cfg["size"], color=C_BODY)
    pf = normal.paragraph_format
    pf.line_spacing = cfg["spacing"]
    pf.space_before = Pt(0)
    pf.space_after = Pt(cfg["after"])
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if cfg["align"] == "both" else WD_ALIGN_PARAGRAPH.LEFT

    for i, name in enumerate(("Heading 1", "Heading 2", "Heading 3")):
        st = doc.styles[name]
        _style_font(st, cfg["en"], cfg["zh"], [cfg["h1"], cfg["h2"], cfg["h3"]][i],
                    bold=True, color=RGBColor.from_string(cfg["hcolor"]))
        p = st.paragraph_format
        p.space_before = Pt(cfg["hbefore"][i]); p.space_after = Pt(cfg["hafter"][i])
        p.keep_with_next = True; p.line_spacing = 1.2
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    from docx.enum.style import WD_STYLE_TYPE
    def add(name, size, color, bold=False, before=0, after=6, left=0, first=0,
            align=WD_ALIGN_PARAGRAPH.LEFT, spacing=1.3):
        st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        _style_font(st, cfg["en"], cfg["zh"], size, bold=bold, color=color)
        f = st.paragraph_format
        f.space_before = Pt(before); f.space_after = Pt(after); f.line_spacing = spacing
        f.left_indent = Cm(left)
        if first:
            f.first_line_indent = Cm(first)
        f.alignment = align
        return st

    add("CoverTitle", 26, RGBColor(0x1F, 0x38, 0x64), bold=True, after=10,
        align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.3)
    add("CoverSub", 14, C_BODY, after=6, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.4)
    add("CoverSmall", 10.5, C_GRAY, after=4, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.4)
    add("CopyHead", 13, C_BODY, bold=True, after=10, spacing=1.2)
    add("CopyLine", 9.5, C_BODY, after=3, spacing=1.4)
    add("CopySmall", 8.5, C_GRAY, after=3, spacing=1.3)
    add("TocHead", 20, RGBColor(0x1F, 0x38, 0x64), bold=True, after=14, spacing=1.2)
    add("Toc1", 11, C_BODY, bold=True, before=4, after=1, spacing=1.1)
    add("Toc2", 10, C_BODY, after=0, left=0.85, spacing=1.0)
    add("Anchor", 9, C_GRAY, after=14, spacing=1.0)
    add("Note", 8.5, RGBColor(0x40, 0x40, 0x40), after=2.5, left=0.95, first=-0.95,
        spacing=1.12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add("Bib", 9, RGBColor(0x22, 0x22, 0x22), after=1.5, left=0.9, first=-0.9, spacing=1.06,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add("BibHead", 13, RGBColor(0x1F, 0x38, 0x64), bold=True, before=18, after=8, spacing=1.2)
    from docx.enum.style import WD_STYLE_TYPE as _ST
    try:
        fts = doc.styles.add_style("FootnoteText", _ST.PARAGRAPH)
        fts.base_style = doc.styles["Normal"]
        fts.font.size = Pt(8.5)
        fts.paragraph_format.space_after = Pt(0)
        fts.paragraph_format.line_spacing = 1.0
        frs = doc.styles.add_style("FootnoteReference", _ST.CHARACTER)
        frs.font.superscript = True
    except Exception:
        pass
    add("FigCap", 9, C_GRAY, before=6, after=12, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.25)

    _first_h1 = next((b for b in book["blocks"] if b["type"] == "h1" and b.get("zh")), None)
    if _first_h1:
        meta.setdefault("title_zh", _first_h1["zh"][:60])
        meta.setdefault("title_en", (_first_h1.get("text") or "")[:120])

    sec1 = doc.sections[0]
    _sect(sec1, cfg["margin"])
    for _ in range(4):
        doc.add_paragraph("", style="CoverSmall")
    doc.add_paragraph(meta.get("title_zh") or meta.get("title", "译文"), style="CoverTitle")
    if meta.get("subtitle"):
        doc.add_paragraph(meta["subtitle"], style="CoverSub")
    _rule = doc.add_paragraph("", style="CoverSmall")
    _border(_rule, color="9FB0C8", sz=6, side="bottom")
    doc.add_paragraph("")
    if meta.get("author"):
        doc.add_paragraph(meta["author"], style="CoverSub")
    if meta.get("author_en") and meta["author_en"] != meta.get("author"):
        doc.add_paragraph(meta["author_en"], style="CoverSmall")
    pub = "　".join(x for x in (meta.get("publisher", ""), meta.get("year", "")) if x)
    if pub:
        doc.add_paragraph(pub, style="CoverSmall")
    for _ in range(4):
        doc.add_paragraph("", style="CoverSmall")
    if meta.get("title_en") and meta["title_en"] != (meta.get("title_zh") or meta.get("title")):
        doc.add_paragraph(meta["title_en"], style="CoverSmall")
    for _ in range(2):
        doc.add_paragraph("", style="CoverSmall")
    imprint = meta.get("imprint", "中译本")
    if imprint:
        doc.add_paragraph(imprint, style="CoverSmall")

    if meta.get("copyright_page", True):
        doc.add_page_break()
        head = doc.add_paragraph("版权信息", style="CopyHead")
        _border(head, color="9FB0C8", sz=4, side="bottom")
        author = " / ".join(dict.fromkeys(
            x for x in (meta.get("author"), meta.get("author_en")) if x))
        rows = [
            ("书名", meta.get("title_zh") or meta.get("title") or ""),
            ("原书名", meta.get("title_en") or ""),
            ("作者", author),
            ("原出版者", meta.get("publisher") or ""),
            ("出版年", meta.get("year") or ""),
            ("译者", "＿＿＿＿"),
            ("出版发行", "＿＿＿＿"),
            ("版次", "＿＿＿＿年＿＿月第 1 版"),
            ("ISBN", "＿＿＿＿"),
        ]
        book_title = rows[0][1]
        seen = set()
        for label, value in rows:
            if label == "原书名" and value == book_title:
                continue
            if not value or (label, value) in seen:
                continue
            seen.add((label, value))
            p = doc.add_paragraph(style="CopyLine")
            r = p.add_run(label + "："); r.bold = True
            p.add_run(str(value))
        doc.add_paragraph("版权所有 · 侵权必究", style="CopyLine")
        doc.add_paragraph("（校样占位；出版信息由出版社核定）", style="CopySmall")

    sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
    _sect(sec2, cfg["margin"])
    sec2.footer.is_linked_to_previous = False
    sec2.header.is_linked_to_previous = False

    for _b in book["blocks"]:
        for _k in ("text", "zh", "caption"):
            if isinstance(_b.get(_k), str):
                _b[_k] = _clean(_b[_k])
    for _k, _v in list(meta.items()):
        if isinstance(_v, str):
            meta[_k] = _clean(_v)

    blocks = [b for b in book["blocks"]
              if b.get("zh") or b["type"] in ("figure", "table", "bib")]
    all_heads = [b for b in blocks if b["type"] in ("h1", "h2")]
    for i, b in enumerate(all_heads):
        b["anchor"] = "_Toc%04d" % (900 + i)
        b["bid"] = 900 + i
    refs, fig_dups = _assign_figure_numbers(blocks)
    recs, endnote_ids, endnote_nos = _collect_notes(blocks)
    st = {
        "notes_total": len(recs), "notes_anchored": 0, "notes_exact": 0,
        "notes_near": 0, "notes_fallback": 0,
        "notes_orphan": 0, "markers_total": 0, "markers_unresolved": 0,
        "markers_endnote": 0, "markers_duplicate": 0, "marker_misses": [], "marker_dups": [], "figures": sum(1 for b in blocks if b["type"] == "figure"),
        "tables": sum(1 for b in blocks if b["type"] == "table"),
        "figure_dups": list(fig_dups), "page_tags": page_tags,
        "copyright_page": bool(meta.get("copyright_page", True)),
        "odd_even_header": bool(meta.get("odd_even_header", True)),
    }
    heads = toc_heads(blocks, meta)

    doc.add_paragraph("目　录", style="TocHead")
    for b in heads:
        p = doc.add_paragraph(style="Toc1" if b["type"] == "h1" else "Toc2")
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        _link(p, b["anchor"], b["zh"], 11 if b["type"] == "h1" else 10,
              color=C_BODY, underline=False)
        if pages and b["anchor"] in pages:
            r = p.add_run("\t" + str(pages[b["anchor"]]))
            r.font.size = Pt(10.5)

    sec3 = doc.add_section(WD_SECTION.NEW_PAGE)
    _sect(sec3, cfg["margin"])
    sec3.footer.is_linked_to_previous = False
    sec3.header.is_linked_to_previous = False
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1")
    sec3._sectPr.append(pg)

    def _folio(par):
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = par.add_run("第 "); run.font.size = Pt(9); run.font.color.rgb = C_GRAY
        _page_field(par)
        run = par.add_run(" 页"); run.font.size = Pt(9); run.font.color.rgb = C_GRAY

    _folio(sec3.footer.paragraphs[0])
    if not meta.get("title_zh"):
        for _b in blocks[:12]:
            if _b["type"] == "h1" and _b.get("zh"):
                meta["title_zh"] = _b["zh"][:40]
                break
    if cfg["header"] == "title":
        title = meta.get("title_zh") or meta.get("title", "")
        if meta.get("odd_even_header", True):
            doc.settings.odd_and_even_pages_header_footer = True
            hp = sec3.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _style_ref_field(hp, "Heading 1", title, 9)
            ehp = sec3.even_page_header
            ehp.is_linked_to_previous = False
            ep = ehp.paragraphs[0]
            ep.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = ep.add_run(title)
            r.font.size = Pt(9); r.font.color.rgb = C_GRAY
            efp = sec3.even_page_footer
            efp.is_linked_to_previous = False
            _folio(efp.paragraphs[0])
        else:
            hp = sec3.header.paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = hp.add_run(title)
            r.font.size = Pt(9); r.font.color.rgb = C_GRAY

    first = True
    started = False
    note_store = footnotes.FootnoteStore()
    endnotes = False

    holder = {"last": None}

    def flush_page(page):
        for r in recs:
            if r["used"] or str(r["page"]) != str(page):
                continue
            if holder["last"] is None:
                continue                    # nothing to anchor to yet; try again later
            fid = note_store.add(r["text"])
            note_store.reference(holder["last"], fid)
            r["used"] = True
            st["notes_fallback"] += 1

    def note_at(no, page, par):
        rec, kind = _note_lookup(recs, no, page)
        st["markers_total"] += 1
        if rec is None:
            if str(no) in endnote_nos:
                st["markers_endnote"] += 1
                return False
            if any(r["used"] and str(r["no"]) == str(no) for r in recs):
                st["markers_duplicate"] += 1
                st["marker_dups"].append("%s@p.%s" % (no, page))
            else:
                st["markers_unresolved"] += 1
                st["marker_misses"].append("%s@p.%s" % (no, page))
            return False
        fid = note_store.add(rec["text"])
        note_store.reference(par, fid)
        rec["used"] = True
        st["notes_anchored"] += 1
        st["notes_exact" if kind == "exact" else "notes_near"] += 1
        return True

    for b in blocks:
        if _is_legal_line(b.get("zh") or b.get("text", "")):
            continue
        t = b["type"]
        if str(b.get("page")) != holder.get("page"):
            flush_page(holder.get("page"))
            holder["page"] = str(b.get("page"))
        if t == "h1":
            plabel = str(b.get("page") or "")
            head = b.get("zh") or b.get("text") or ""
            endnotes = bool(b.get("endnote")) or bool(NOTE_HEAD.match(head))
            is_part = bool(_re.match(r"^(?:第[一二三四五六七八九十百]+部分|PART\b)", head, _re.I))
            brk = bool(cfg.get("chapter_break") and started
                       and (plabel.isdigit() or not plabel))
            spacer = not first and cfg["rule"] and not is_part
            if spacer:
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(24)
                sp.paragraph_format.space_after = Pt(0)
                sp.paragraph_format.line_spacing = Pt(1)
                _border(sp)
                if brk:
                    sp.paragraph_format.page_break_before = True
            first = False
            notes_open = False
            p = doc.add_paragraph(style="Heading 1")
            _add_marked(p, _clean(b["zh"]),
                        on_note=lambda no, pg=b.get("page"), par=p: note_at(no, pg, par),
                        refs=refs, size=cfg["h1"], en=cfg["en"], zh=cfg["zh"])
            if brk and not spacer:
                p.paragraph_format.page_break_before = True
            if is_part:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(200)
                p.paragraph_format.space_after = Pt(18)
                for r in p.runs:
                    r.font.size = Pt(20)
            p.paragraph_format.keep_with_next = True
            _bookmark(p, b["anchor"], b["bid"])
        elif t == "h2":
            endnotes = bool(b.get("endnote")) or bool(
                NOTE_HEAD.match((b.get("text") or "").strip())
                or NOTE_HEAD.match((b.get("zh") or "").strip()))
            p = doc.add_paragraph(style="Heading 2")
            _add_marked(p, b["zh"],
                        on_note=lambda no, pg=b.get("page"), par=p: note_at(no, pg, par),
                        refs=refs, size=cfg["h2"], en=cfg["en"], zh=cfg["zh"])
            if cfg.get("chapter_break") and started and re.fullmatch(
                    r"[ivxlcdm]{1,7}", str(b.get("page") or "")):
                p.paragraph_format.page_break_before = True
            p.paragraph_format.keep_with_next = True
            _bookmark(p, b["anchor"], b["bid"])
        elif t == "h3":
            p = doc.add_paragraph(style="Heading 3")
            _add_marked(p, b["zh"],
                        on_note=lambda no, pg=b.get("page"), par=p: note_at(no, pg, par),
                        refs=refs, size=cfg["h3"], en=cfg["en"], zh=cfg["zh"])
            p.paragraph_format.keep_with_next = True
        elif t == "byline":
            b.setdefault("zh", b["text"])
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(16)
            r = p.add_run(b["zh"])
            r.italic = True
            r.font.size = Pt(11)
        elif t == "table":
            if not b.get("image"):
                continue
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(6)
            pic.paragraph_format.space_after = Pt(4)
            pic.add_run().add_picture(b["image"], width=Cm(16.2))
            cap = doc.add_paragraph(b.get("caption") or "", style="FigCap")
            _bookmark(cap, b["anchor"], b["bid"])
        elif t == "figure":
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.space_before = Pt(6)
            pic.paragraph_format.space_after = Pt(4)
            pic.add_run().add_picture(b["image"], width=Cm(16.2))
            cap = doc.add_paragraph(b.get("caption") or "", style="FigCap")
            _bookmark(cap, b["anchor"], b["bid"])
        elif t == "bib":
            b["zh"] = b["text"]          # references stay in the original language
            _z = b["zh"].strip()
            _heads = []
            for _h in blocks:
                if _h["type"] == "h1":
                    _heads += [(_h.get("zh") or "")[:24], (_h.get("text") or "")[:24]]
            if any(h and _z.startswith(h) for h in _heads if h):
                continue
            if re.match(r"^(https?://|www\.)", _z) or _z[:1].islower() or "http" in _z[:70]:
                continue                      # orphan fragment of a reference entry
            _bp = doc.add_paragraph(style="Bib")
            _add_marked(_bp, _bib_marked(b["zh"]))
        elif t == "note":
            if id(b) not in endnote_ids:
                continue
            for no, text in _note_parts(b.get("text", ""), b.get("zh") or ""):
                np = doc.add_paragraph(style="Note")
                if page_tags:
                    _page_tag_run(np, b.get("page"), b.get("page_end"))
                _add_marked(np, (("%s. " % no) if no else "") + _clean(text))
        else:
            p = doc.add_paragraph(style="Normal")
            if page_tags:
                _page_tag_run(p, b.get("page"), b.get("page_end"))
            _add_marked(p, b["zh"],
                        on_note=lambda no, pg=b.get("page"), par=p: note_at(no, pg, par),
                        refs=refs, size=cfg["size"], en=cfg["en"], zh=cfg["zh"])
            holder["last"] = p
            if (b.get("zh") or "").lstrip()[:1] in "●•·◦▪":
                # keep the source bullet glyph but hang it, so wrapped lines
                # align with the text instead of falling back to the margin
                p.paragraph_format.left_indent = Pt(cfg["size"] * 1.25)
                p.paragraph_format.first_line_indent = Pt(-cfg["size"] * 1.25)
                p.paragraph_format.space_after = Pt(1.5)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif b.get("align") == "right":
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif cfg["indent"]:
                _first_line_chars(p)
            if cfg["align"] == "both" and b.get("align") != "right":
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        started = True

    remaining = [r for r in recs if not r["used"]]
    if remaining:
        if holder["last"] is not None:
            for r in remaining:
                fid = note_store.add(r["text"])
                note_store.reference(holder["last"], fid)
                r["used"] = True
                st["notes_fallback"] += 1
        else:
            doc.add_paragraph("脚注", style="Heading 2")
            for r in remaining:
                np = doc.add_paragraph(style="Note")
                _add_marked(np, r["text"])
                st["notes_orphan"] += 1
    note_store.attach(doc)
    if stats is not None:
        stats.clear()
        stats.update(st)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
