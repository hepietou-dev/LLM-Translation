"""DOCX input: convert to PDF, then reuse the PDF pipeline.

Word already carries the chapter structure as paragraph styles and the layout
as real pages, so the cheapest faithful route is a headless conversion to PDF
(LibreOffice, the bundled kit, or Word/WPS over COM) followed by the normal
PDF pipeline.  The style outline is read separately and stamped onto the
ingested blocks, which beats guessing heading levels from font size.
"""
from __future__ import annotations

import base64
import re
import shutil
import subprocess
import time
from pathlib import Path

from .procutil import no_window

HEAD_NAME = re.compile(r"^(?:heading|标题)\s*([1-9])$", re.I)
SKIP_NAME = re.compile(
    r"(?i)^(?:title|subtitle|toc\s*\d|contents|caption|footnote|"
    r"header|footer|page\s*number|书目|标题|副标题|题注|脚注|页眉|页脚)$")
CONTENTS = re.compile(r"^(?:contents|table of contents|目\s*录)$", re.I)


def _level_of(paragraph) -> int | None:
    """1-3 for Heading 1/2/3 styles (outline level or style id/name)."""
    try:
        ppr = paragraph._p.pPr
        if ppr is not None:
            lvl = ppr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}outlineLvl")
            if lvl is not None:
                v = lvl.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val")
                if v is not None and int(v) < 3:
                    return int(v) + 1
        style = paragraph.style
        while style is not None:
            sid = (style.style_id or "").lower()
            name = (style.name or "").strip()
            m = HEAD_NAME.match(name) or HEAD_NAME.match(sid)
            if m:
                return min(3, int(m.group(1)))
            if SKIP_NAME.match(name) or SKIP_NAME.match(sid):
                return None
            style = style.base_style
    except Exception:                                     # noqa: BLE001
        return None
    return None


def headings(path: Path) -> list[dict]:
    """(level, text) for every styled heading, in document order."""
    from docx import Document

    doc = Document(str(path))
    out: list[dict] = []
    for p in doc.paragraphs:
        text = " ".join((p.text or "").split())
        if not text or len(text) > 200 or CONTENTS.match(text):
            continue
        lvl = _level_of(p)
        if lvl:
            out.append({"level": lvl, "text": text})
    return out


def _split_title(title: str) -> tuple[str, str]:
    m = re.match(r"^(.+?)\s*[:：]\s+(.+)$", title.strip())
    if m and len(m.group(1)) > 3:
        return m.group(1).strip(), m.group(2).strip()
    return title.strip(), ""


def core_meta(path: Path) -> dict:
    """Title/author/year/publisher from the DOCX package properties."""
    from docx import Document
    import zipfile

    info = {"title": "", "subtitle": "", "author": "", "publisher": "", "year": ""}
    try:
        doc = Document(str(path))
        cp = doc.core_properties
        title = (cp.title or "").strip()
        author = (cp.author or "").strip()
        if not title:
            for p in doc.paragraphs:
                if (p.style is not None and (p.style.name or "").lower() == "title"
                        and (p.text or "").strip()):
                    title = " ".join(p.text.split())
                    break
        if not author:
            author = (cp.last_modified_by or "").strip()
        if cp.created:
            info["year"] = str(cp.created.year)
    except Exception:                                     # noqa: BLE001
        title, author = "", ""
    if title:
        info["title"], info["subtitle"] = _split_title(title)
    info["author"] = re.sub(r"\s+", " ", author).strip()
    try:
        with zipfile.ZipFile(path) as z:
            app = z.read("docProps/app.xml").decode("utf-8", "replace")
        m = re.search(r"<Company>(.*?)</Company>", app, re.S | re.I)
        if m:
            info["publisher"] = re.sub(r"\s+", " ", m.group(1)).strip()
    except (KeyError, OSError, ValueError):
        pass
    return info


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", (s or "").casefold())


def drop_front_matter(blocks: list[dict], meta: dict) -> int:
    """The cover is rebuilt from metadata, so drop a duplicated leading title/author."""
    keys = {_norm(meta.get(k) or "") for k in ("title", "subtitle", "author")}
    keys = {k for k in keys if len(k) >= 4}
    if not keys:
        return 0
    out, leading, dropped = [], True, 0
    for b in blocks:
        if leading:
            t = _norm(b.get("text") or "")
            if t and t in keys:
                dropped += 1
                continue
            if b["type"] in ("figure", "table"):
                leading = False
            elif t:
                leading = False
        out.append(b)
    if dropped:
        blocks[:] = out
        for n, b in enumerate(blocks):
            b["id"] = "b%04d" % n
    return dropped


def _cut_prefix(text: str, head: str) -> str | None:
    """Text after ``head`` at the start of ``text`` (spacing-insensitive)."""
    if text[:len(head)].casefold() == head.casefold():
        return text[len(head):].lstrip(" \u3000:：")
    pos = 0
    for word in head.split():
        m = re.search(r"[^\S\r\n]*" + re.escape(word), text[pos:], re.I)
        if not m:
            return None
        pos += m.end()
    return text[pos:].lstrip(" \u3000:：")


def apply_headings(blocks: list[dict], heads: list[dict]) -> int:
    """Stamp Word's outline onto the ingested blocks, in reading order."""
    kind = {1: "h1", 2: "h2", 3: "h3"}
    applied, j = 0, 0
    for h in heads:
        key = _norm(h["text"])
        if len(key) < 2:
            continue
        hit = next((i for i in range(j, len(blocks))
                    if blocks[i]["type"] not in ("figure", "table")
                    and _norm(blocks[i].get("text")) == key), None)
        rest = None
        if hit is None:
            # the heading shares a PDF block with its first paragraph
            for i in range(j, len(blocks)):
                b = blocks[i]
                if b["type"] in ("figure", "table"):
                    continue
                nt = _norm(b.get("text") or "")
                if len(nt) > len(key) and nt.startswith(key):
                    cut = _cut_prefix(b.get("text") or "", h["text"])
                    if cut is None or len(_norm(cut)) < 20:
                        continue
                    hit, rest = i, cut
                    break
        if hit is None:
            continue
        b = blocks[hit]
        b["type"] = kind[h["level"]]
        if rest:
            b["text"] = " ".join(h["text"].split())
            blocks.insert(hit + 1, {"page": b.get("page"), "type": "p", "text": rest})
            j = hit + 2
        else:
            j = hit + 1
        applied += 1
    return applied


def _run(cmd: list[str], stop=None, timeout: float = 1800) -> bool:
    if stop is None:
        p = subprocess.run(cmd, capture_output=True, timeout=timeout, **no_window())
        return p.returncode == 0
    p = subprocess.Popen(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, **no_window())
    started = time.monotonic()
    while p.poll() is None:
        if stop.is_set() or time.monotonic() - started > timeout:
            p.terminate()
            try:
                p.wait(5)
            except Exception:                             # noqa: BLE001
                p.kill()
            return False
        time.sleep(0.2)
    return p.returncode == 0


def _word_com(src: Path, out: Path, stop=None) -> bool:
    exe = shutil.which("powershell") or shutil.which("powershell.exe")
    if not exe:
        return False
    script = r'''
$ErrorActionPreference = "SilentlyContinue"
$src = '%s'
$out = '%s'
$existing = $true
try { $null = [Runtime.InteropServices.Marshal]::GetActiveObject("Word.Application") }
catch { $existing = $false }
$app = $null
foreach ($id in @("Word.Application", "KWPS.Application")) {
    try { $app = New-Object -ComObject $id; break } catch { }
}
if ($null -eq $app) { exit 3 }
$doc = $null
try {
    if (-not $existing) { try { $app.Visible = $false } catch { } }
    try { $app.DisplayAlerts = 0 } catch { }
    try { $doc = $app.Documents.Open($src, $false, $true, $false, "", "", $true, "", "", 0, 0, $false) }
    catch { $doc = $app.Documents.Open($src) }
    if ($null -eq $doc) { exit 4 }
    try { $doc.ExportAsFixedFormat($out, 17) }
    catch {
        try { $doc.SaveAs2($out, 17) }
        catch { $doc.SaveAs([ref]$out, [ref]17) }
    }
} finally {
    if ($null -ne $doc) { try { $doc.Close(0) } catch { } }
    if (-not $existing) { try { $app.Quit() } catch { } }
}
''' % (str(src).replace("'", "''"), str(out).replace("'", "''"))
    enc = base64.b64encode(script.encode("utf-16-le")).decode("ascii")
    _run([exe, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
          "-WindowStyle", "Hidden", "-EncodedCommand", enc], stop=stop)
    return out.exists() and out.stat().st_size > 0


def convert(path: Path, out_pdf: Path, *, soffice: str | None = None,
            node: str | None = None, cli: str | None = None,
            log=print, stop=None) -> bool:
    """DOCX -> PDF through the first available headless converter."""
    path, out_pdf = Path(path), Path(out_pdf)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    out_pdf.unlink(missing_ok=True)

    if node and cli and Path(node).exists() and Path(cli).exists():
        if _run([node, cli, "convert", "--input", str(path), "--output", str(out_pdf)], stop=stop):
            if out_pdf.exists() and out_pdf.stat().st_size:
                return True
    if soffice:
        _run([soffice, "--headless", "--norestore", "--convert-to", "pdf",
              "--outdir", str(out_pdf.parent), str(path)], stop=stop)
        produced = out_pdf.parent / (path.stem + ".pdf")
        if produced.exists():
            produced.replace(out_pdf)
            return True
    if _word_com(path, out_pdf, stop=stop):
        return True
    log("docx: 未找到可用的转换器（需要 Microsoft Word、WPS 或 LibreOffice）")
    return False
