"""Fast, offline unit tests (no API, no large fixtures)."""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]

import _bootstrap  # noqa: E402,F401  (side effect: sys.path)

import pymupdf  # noqa: E402
from PIL import Image  # noqa: E402

from llmtr import ingest, ocr, pipeline  # noqa: E402


def _page(tmp: Path, w: float, h: float, raster: bool):
    doc = pymupdf.open()
    page = doc.new_page(width=w, height=h)
    if raster:
        img = tmp / "page.png"
        Image.new("RGB", (600, 420), "white").save(img)
        page.insert_image(page.rect, filename=str(img))
    return page


class SpreadDetection(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_landscape_raster_is_a_spread(self):
        page = _page(self.tmp, 1600, 1100, raster=True)
        self.assertTrue(ocr.is_spread(page))
        self.assertEqual(len(ocr.halves(page)), 2)

    def test_portrait_scan_is_not_a_spread(self):
        page = _page(self.tmp, 1000, 1600, raster=True)
        self.assertFalse(ocr.is_spread(page))

    def test_landscape_without_raster_is_not_a_spread(self):
        page = _page(self.tmp, 1600, 1100, raster=False)
        self.assertFalse(ocr.is_spread(page))

    def test_modes_override(self):
        portrait = _page(self.tmp, 1000, 1600, raster=True)
        self.assertFalse(ocr.is_spread(portrait, "off"))
        self.assertTrue(ocr.is_spread(portrait, "force"))


class MarkedLayout(unittest.TestCase):
    def test_parse_marked_keeps_structure(self):
        label, items = ocr.parse_marked(
            "[PAGE] 42\n[H2] ESTABLISHING POWER\n[P] first paragraph\n"
            "[N] 12 a footnote that\nwraps onto a second line")
        self.assertEqual(label, "42")
        self.assertEqual(items[0], ("h2", "ESTABLISHING POWER"))
        self.assertEqual(items[1], ("p", "first paragraph"))
        self.assertEqual(items[2], ("note", "12 a footnote that wraps onto a second line"))

    def test_page_marker_must_be_a_number(self):
        label, _items = ocr.parse_marked("[PAGE] ix\n[P] text")
        self.assertEqual(label, "ix")
        label, _items = ocr.parse_marked("[PAGE] ?\n[P] text")
        self.assertIsNone(label)


class SpreadLabels(unittest.TestCase):
    def test_blanks_are_filled_between_printed_numbers(self):
        layout = [(None, []), ("20", []), ("21", []), (None, []), ("23", [])]
        self.assertEqual(ingest._layout_labels(layout), ["", "20", "21", "22", "23"])

    def test_misread_digit_between_romans_becomes_roman(self):
        layout = [("vi", []), ("6", []), ("vii", [])]
        self.assertEqual(ingest._layout_labels(layout), ["vi", "vi", "vii"])


class LayoutIngest(unittest.TestCase):
    """A spread page must yield two printed pages pointing at one PDF sheet."""

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.TemporaryDirectory()
        cls.pdf = Path(cls._tmp.name) / "spread.pdf"
        doc = pymupdf.open()
        for i in range(3):
            doc.new_page(width=1000 if i == 0 else 1600, height=1600 if i == 0 else 1100)
        doc.save(str(cls.pdf))
        doc.close()

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_labels_and_page_map(self):
        layout = [
            {"label": None, "items": [("p", "Title page.")], "src": 0, "half": None},
            {"label": "20", "items": [("p", "Left column paragraph."),
                                       ("note", "9 a left footnote")], "src": 1, "half": 0},
            {"label": "21", "items": [("p", "Right column paragraph.")], "src": 1, "half": 1},
            {"label": "22", "items": [("p", "Next spread left.")], "src": 2, "half": 0},
            {"label": "23", "items": [("p", "Next spread right.")], "src": 2, "half": 1},
        ]
        book = ingest.ingest(self.pdf, layout=layout)
        self.assertTrue(book["ok"])
        self.assertEqual(book["labels"], ["", "20", "21", "22", "23"])
        self.assertEqual(book["page_map"], {"": 0, "20": 1, "21": 1, "22": 2, "23": 2})
        notes = [b for b in book["blocks"] if b["type"] == "note"]
        self.assertEqual([b["page"] for b in notes], ["20"])
        pages = {b["page"] for b in book["blocks"]}
        self.assertEqual(pages, {"20", "21", "22", "23"})   # unnumbered cover dropped

    def test_page_index_uses_the_map(self):
        page_map = {"20": 1, "21": 1}
        self.assertEqual(ingest.page_index("21", ["20", "21"], page_map), 1)
        self.assertEqual(ingest.labels_for_src(["20", "21"], page_map, 1), ["20", "21"])


class LayoutPagesSplitting(unittest.TestCase):
    """layout_pages must render each half separately and keep the source page."""

    def test_spread_is_transcribed_half_by_half(self):
        from unittest import mock

        from llmtr import vision

        def fake(_client, img, max_side=1600, stop=None):
            half = Path(img).stem.split("_")[-1]
            return "[PAGE] %s\n[P] %s half text." % (
                {"0": "20", "1": "21"}[half], "left" if half == "0" else "right")

        with tempfile.TemporaryDirectory() as d:
            tmp = Path(d)
            pdf = tmp / "spread.pdf"
            doc = pymupdf.open()
            page = doc.new_page(width=1600, height=1100)
            img = tmp / "art.png"
            Image.new("RGB", (600, 420), "white").save(img)
            page.insert_image(page.rect, filename=str(img))
            doc.save(str(pdf))
            doc.close()

            with mock.patch.object(vision, "transcribe_layout", side_effect=fake):
                out = ocr.layout_pages(pdf, object(), {"layout_dpi": 72, "max_side": 400})

        self.assertEqual([e["label"] for e in out], ["20", "21"])
        self.assertEqual([e["half"] for e in out], [0, 1])
        self.assertEqual([e["src"] for e in out], [0, 0])
        self.assertEqual(out[0]["items"], [("p", "left half text.")])
        self.assertEqual(out[1]["items"], [("p", "right half text.")])


class ChapterLabelTitles(unittest.TestCase):
    """'Chapter One' + title is one chapter; unlabelled headings under it are sections."""

    def test_label_and_title_merge(self):
        blocks = [
            {"page": "17", "type": "h1", "text": "Chapter One"},
            {"page": "17", "type": "h1", "text": "Political Foundations of Class Power"},
            {"page": "17", "type": "p", "text": "Body text starts here."},
        ]
        out = ingest._merge_label_titles([dict(b) for b in blocks])
        self.assertEqual(out[0]["text"],
                         "Chapter One Political Foundations of Class Power")
        self.assertEqual(len(out), 2)

    def test_unlabelled_heading_after_chapters_is_a_section(self):
        blocks = [
            {"page": "1", "type": "h1", "text": "PART ONE"},
            {"page": "1", "type": "h1", "text": "Building Socialism"},
            {"page": "3", "type": "h1", "text": "Chapter One"},
            {"page": "3", "type": "h1", "text": "Political Foundations"},
            {"page": "5", "type": "h2", "text": "Direct subsection"},
            {"page": "7", "type": "h1", "text": "THE PARTY ORGANIZATION"},
            {"page": "8", "type": "h2", "text": "Recruitment"},
            {"page": "20", "type": "h1", "text": "Chapter Two"},
            {"page": "20", "type": "h1", "text": "Cultural Foundations"},
            {"page": "22", "type": "h2", "text": "Another subsection"},
        ]
        blocks = ingest._merge_label_titles([dict(b) for b in blocks])
        out = ingest._demote_section_h1s(blocks)
        kinds = [(b["type"], b["text"]) for b in out]
        self.assertIn(("h2", "THE PARTY ORGANIZATION"), kinds)
        self.assertIn(("h3", "Recruitment"), kinds)
        self.assertIn(("h2", "Direct subsection"), kinds)
        self.assertIn(("h1", "Chapter Two Cultural Foundations"), kinds)
        self.assertIn(("h2", "Another subsection"), kinds)

    def test_numbered_chapter_keeps_h1_when_parts_exist(self):
        blocks = [
            {"page": "6", "type": "h1", "text": "PART I METHODOLOGICAL APPROACHES"},
            {"page": "7", "type": "h1", "text": "1. Philosophical analysis"},
            {"page": "7", "type": "h2", "text": "1.1 INTRODUCTION"},
        ]
        out = ingest._demote_section_h1s([dict(b) for b in blocks])
        self.assertEqual([b["type"] for b in out], ["h1", "h1", "h2"])


class DuplicateChapterTitles(unittest.TestCase):
    """A running head promoted to h1 must not become a second chapter in the TOC."""

    def test_running_head_heading_is_dropped(self):
        from llmtr import clean
        blocks = [
            {"page": "xi", "type": "h1", "text": "EDITOR'S INTRODUCTION", "zh": "编者导言"},
            {"page": "xii", "type": "p", "text": "Text.", "zh": "正文。" * 20},
            {"page": "xvii", "type": "h1", "text": "EDITOR'S INTRODUCTION", "zh": "编者导言"},
            {"page": "xviii", "type": "p", "text": "More.", "zh": "更多。" * 20},
        ]
        out = clean.headings([dict(b) for b in blocks])
        self.assertEqual([b["type"] for b in out].count("h1"), 1)

    def test_same_title_in_a_later_part_is_kept(self):
        from llmtr import clean
        blocks = [
            {"page": "20", "type": "h1", "text": "INTRODUCTION", "zh": "导论"},
            {"page": "200", "type": "h1", "text": "INTRODUCTION", "zh": "导论"},
        ]
        out = clean.headings([dict(b) for b in blocks])
        self.assertEqual([b["type"] for b in out].count("h1"), 2)


class PublicationStructure(unittest.TestCase):
    """Markers, endnote headings, run-in h3 and TOC filters."""

    def test_superscript_note_marker_is_kept(self):
        line = {"spans": [
            {"text": "has none.", "flags": 4},
            {"text": "1", "flags": 5},
            {"text": " As such", "flags": 4},
        ]}
        self.assertEqual(ingest._line_text(line), "has none.\u27e61\u27e7 As such")

    def test_ordinary_digit_span_is_left_alone(self):
        line = {"spans": [{"text": "in 1952, at 71.", "flags": 4}]}
        self.assertEqual(ingest._line_text(line), "in 1952, at 71.")

    def test_runin_h3_splits_heading_from_body(self):
        self.assertEqual(
            ingest.split_runin_h3(
                "6.3.1 The Conjunctural On the conjunctural and most immediate level"),
            ("6.3.1 The Conjunctural", "On the conjunctural and most immediate level"))
        self.assertEqual(
            ingest.split_runin_h3(
                "6.3.3 Science and Tradition Finally we may briefly consider"),
            ("6.3.3 Science and Tradition", "Finally we may briefly consider"))

    def test_runin_h3_ignores_plain_prose(self):
        self.assertIsNone(ingest.split_runin_h3("6.3.1 is discussed in the next chapter"))
        self.assertIsNone(ingest.split_runin_h3("Section 6.3.1 explains this."))

    def test_note_parts_do_not_double_the_period(self):
        from llmtr import build
        self.assertEqual(build._note_parts("10 . This can be explained by clothes.",
                                           "【测试】10 . This can be explained by clothes."),
                         [("10", "This can be explained by clothes.")])
        self.assertEqual(build._note_parts("3. See Strauss (1952, 71).", "3. 参见 Strauss（1952, 71）。"),
                         [("3", "参见 Strauss（1952, 71）。")])

    def test_toc_drops_repeated_note_and_reference_headings(self):
        from llmtr import build
        blocks = [
            {"type": "h1", "page": "1", "text": "Introduction", "zh": "导论"},
            {"type": "h2", "page": "4", "text": "NOTE", "zh": "注释"},
            {"type": "h2", "page": "5", "text": "References", "zh": "参考文献"},
            {"type": "h1", "page": "5", "text": "References", "zh": "参考文献"},
            {"type": "h1", "page": "6", "text": "PART I", "zh": "第一部分"},
            {"type": "h1", "page": "6", "text": "PART I", "zh": "第一部分"},
        ]
        got = [b["zh"] for b in build.toc_heads(blocks, {"author": "Nederman"})]
        self.assertEqual(got, ["导论", "第一部分"])

    def test_marker_run_is_superscript(self):
        from docx import Document

        from llmtr import build
        doc = Document()
        p = doc.add_paragraph()
        build._add_marked(p, "见注\u27e63\u27e7。")
        self.assertEqual([r.text for r in p.runs], ["见注", "3", "。"])
        self.assertTrue(p.runs[1].font.superscript)

    def test_invented_middle_initial_is_dropped(self):
        from llmtr import clean
        blocks = [
            {"page": "1", "type": "p", "text": "Kelvin Knight is a fellow.",
             "zh": "凯文·J. 奈特是研究员。"},
            {"page": "1", "type": "p", "text": "John W. Wetzel is a professor.",
             "zh": "约翰·W. 韦策尔是教授。"},
        ]
        out = clean.finalize(blocks, None)
        self.assertIn("凯文·奈特", out[0]["zh"])
        self.assertIn("约翰·W. 韦策尔", out[1]["zh"])

    def test_byline_follows_contributor_bio(self):
        from llmtr import clean
        blocks = [
            {"page": "xiv", "type": "p", "text": "Takamichi Sakurai currently teaches at Goethe University.",
             "zh": "樱井孝道（Takamichi Sakurai）目前在法兰克福歌德大学任教。"},
            {"page": "42", "type": "byline", "text": "Takamichi Sakurai",
             "zh": "高道坂井（Takamichi Sakurai）"},
        ]
        out = clean.finalize(blocks, None)
        self.assertEqual(out[1]["zh"], "樱井孝道（Takamichi Sakurai）")

    def test_glued_edited_volume_entry_is_split(self):
        from llmtr import clean
        blocks = [{
            "page": "39", "type": "bib",
            "text": "Anderson, Perry. 1992. “Components of the National Culture,” in English "
                    "Questions, 48–104. London: Verso. Barker, Ernest, ed. The Politics of "
                    "Aristotle. Oxford: Oxford University Press."}]
        out = clean.run(blocks)
        self.assertEqual(len(out), 2)
        self.assertTrue(out[1]["text"].startswith("Barker, Ernest, ed. The Politics"))

    def test_glued_author_year_entry_is_split(self):
        from llmtr import clean
        blocks = [{
            "page": "60", "type": "bib",
            "text": "Akkerman, Tjiske and Siep Stuurman, eds. 1998. Perspectives on Feminist "
                    "Political Thought. London: Routledge. Albret, Jeanne d’. 1572. Reprinted "
                    "1970. Mémoires et Poésies. Geneva: Slatkine Reprints."}]
        out = clean.run(blocks)
        self.assertEqual(len(out), 2)
        self.assertTrue(out[1]["text"].startswith("Albret, Jeanne d’"))

    def test_reference_missing_period_is_split(self):
        from llmtr import clean
        blocks = [{
            "page": "28", "type": "bib",
            "text": "Strauss, Leo. 1939. “The Spirit of Sparta.” Social Research 6: 502–36 "
                    "Strauss, Leo. 1945. “Farabi’s Plato,” Louis Ginzberg: Jubilee Volume. "
                    "357–93. New York: American Academy for Jewish Research."}]
        out = clean.run(blocks)
        self.assertEqual(len(out), 2)
        self.assertTrue(out[1]["text"].startswith("Strauss, Leo. 1945."))

    def test_note_marker_drops_space_before_chinese(self):
        from docx import Document

        from llmtr import build
        doc = Document()
        p = doc.add_paragraph()
        build._add_marked(p, "见注\u27e61\u27e7 正文。")
        self.assertEqual([r.text for r in p.runs], ["见注", "1", "正文。"])

    def test_heading_style_overrides_theme_fonts(self):
        from docx import Document
        from docx.oxml.ns import qn

        from llmtr import build
        doc = Document()
        build._style_font(doc.styles["Heading 1"], "Cambria", "宋体", 20, bold=True)
        rf = doc.styles["Heading 1"].element.rPr.rFonts
        self.assertIsNone(rf.get(qn("w:asciiTheme")))
        self.assertIsNone(rf.get(qn("w:eastAsiaTheme")))
        self.assertEqual(rf.get(qn("w:eastAsia")), "宋体")

    def test_right_aligned_signature_block(self):
        import tempfile
        from pathlib import Path

        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        from llmtr import build
        book = {"blocks": [
            {"id": "s1", "type": "p", "zh": "卡里·J. 内德曼",
             "text": "Cary J. Nederman", "align": "right"},
        ], "meta": {"title": "T"}}
        out = Path(tempfile.gettempdir()) / "llmtr_sig_test.docx"
        build.build_doc(book, "print", out, book["meta"])
        doc = Document(str(out))
        p = next(p for p in doc.paragraphs if p.text.strip() == "卡里·J. 内德曼")
        self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertIsNone(p.paragraph_format.first_line_indent)

    def test_cover_has_no_tool_watermark(self):
        import tempfile
        from pathlib import Path

        from docx import Document

        from llmtr import build
        book = {"blocks": [], "meta": {"title": "T"}}
        out = Path(tempfile.gettempdir()) / "llmtr_cover_test.docx"
        build.build_doc(book, "print", out, book["meta"])
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        self.assertNotIn("LLM-Translation", text)
        self.assertIn("中译本", text)

    def test_bullet_paragraph_hangs_glyph(self):
        import tempfile
        from pathlib import Path

        from docx import Document
        from docx.oxml.ns import qn

        from llmtr import build
        book = {"blocks": [{
            "id": "p1", "type": "p", "page": "1",
            "zh": "● 概念清晰既需要能力也需要实践。",
            "text": "Conceptual clarity requires both ability and practice.",
        }], "meta": {"title": "T"}}
        out = Path(tempfile.gettempdir()) / "llmtr_bullet_test.docx"
        build.build_doc(book, "print", out, book["meta"])
        p = next(p for p in Document(str(out)).paragraphs
                 if p.text.startswith("●"))
        self.assertGreater(p.paragraph_format.left_indent.pt, 0)
        self.assertLess(p.paragraph_format.first_line_indent.pt, 0)
        ind = p._p.pPr.find(qn("w:ind"))
        self.assertIsNone(ind.get(qn("w:firstLineChars")))

    def test_bib_titles_are_italicized(self):
        from llmtr import build
        cases = [
            ("Kogelmann, Brian. 2018. “Enough and as Good.” American Journal of Political "
             "Science 62: 682–94.",
             "Kogelmann, Brian. 2018. “Enough and as Good.” §American Journal of Political "
             "Science§ 62: 682–94."),
            ("Blau, Adrian. 2016. “Reason and the Passions,” in The Oxford Handbook of "
             "Hobbes, 195–220. Ed. A.P. Martinich. Oxford: Oxford University Press.",
             "Blau, Adrian. 2016. “Reason and the Passions,” in §The Oxford Handbook of "
             "Hobbes§, 195–220. Ed. A.P. Martinich. Oxford: Oxford University Press."),
            ("Portis, Edward Bryan. 2008. Reconstructing the Classics, 3rd edn. "
             "Washington, DC: CQ Press.",
             "Portis, Edward Bryan. 2008. §Reconstructing the Classics§, 3rd edn. "
             "Washington, DC: CQ Press."),
            ("Gray, John. 1996. Mill On Liberty: A Defence. 2nd edn. London: Routledge.",
             "Gray, John. 1996. §Mill On Liberty: A Defence§. 2nd edn. London: Routledge."),
            ("Albret, Jeanne d’. 1572. Reprinted 1970. Mémoires et Poésies. "
             "Geneva: Slatkine Reprints.",
             "Albret, Jeanne d’. 1572. Reprinted 1970. §Mémoires et Poésies§. "
             "Geneva: Slatkine Reprints."),
            ("Barker, Ernest, ed. The Politics of Aristotle. Oxford: Oxford University Press.",
             "Barker, Ernest, ed. §The Politics of Aristotle§. Oxford: Oxford University Press."),
            ("Major, Rafael, ed. 2013. Leo Strauss’s Defense of the Philosophic Life: "
             "Reading “What Is Political Philosophy?” Chicago: University of Chicago Press.",
             "Major, Rafael, ed. 2013. §Leo Strauss’s Defense of the Philosophic Life: "
             "Reading “What Is Political Philosophy?”§ Chicago: University of Chicago Press."),
            ("Pocock, J.G.A. 1962. “The History of Political Thought,” in Politics, "
             "Philosophy and Society (second series). Ed. Peter Laslett, 183–202. "
             "Oxford: Basil Blackwell.",
             "Pocock, J.G.A. 1962. “The History of Political Thought,” in §Politics, "
             "Philosophy and Society (second series)§. Ed. Peter Laslett, 183–202. "
             "Oxford: Basil Blackwell."),
            ("Gibbons, Rachel. 1996. “Isabeau of Bavaria.” Transactions of the Royal "
             "Historical Society Ser. 6, 6: 51–73.",
             "Gibbons, Rachel. 1996. “Isabeau of Bavaria.” §Transactions of the Royal "
             "Historical Society§ Ser. 6, 6: 51–73."),
        ]
        for src, expect in cases:
            self.assertEqual(build._bib_marked(src).replace("*", "§"), expect)
        self.assertEqual(build._bib_marked("Alexander　亚历山大　2, 40"),
                         "Alexander　亚历山大　2, 40")

    def test_bib_reference_docx_runs_have_italics(self):
        import tempfile
        from pathlib import Path

        from docx import Document

        from llmtr import build
        book = {"blocks": [{
            "id": "b1", "type": "bib", "page": "1",
            "text": "Berlin, Isaiah. 1969. Four Essays on Liberty. Oxford: Oxford "
                    "University Press.",
        }], "meta": {"title": "T"}}
        out = Path(tempfile.gettempdir()) / "llmtr_bib_test.docx"
        build.build_doc(book, "print", out, book["meta"])
        p = next(p for p in Document(str(out)).paragraphs
                 if p.text.startswith("Berlin, Isaiah"))
        italics = [r.text for r in p.runs if r.italic]
        self.assertEqual(italics, ["Four Essays on Liberty"])

    def test_untranslated_terms_are_flagged(self):
        from llmtr import proofread
        self.assertEqual(proofread.suspect_terms("康德的 cosmopolitanism 可以说是白人的"),
                         ["cosmopolitanism"])
        self.assertEqual(proofread.suspect_terms("发表于Journal of Political Philosophy。"), [])
        self.assertEqual(proofread.suspect_terms("他所说的 virtù 与 necessità"), [])
        self.assertEqual(proofread.suspect_terms("见 https://example.com/page。"), [])

    def test_citations_and_titles_are_not_stray_terms(self):
        from llmtr import proofread
        # names joined by "and", journal titles, Latin titles, Greek catalogues
        self.assertEqual(proofread.suspect_terms(
            "参见 Smith (2007)、Zuckert and Zuckert (2014)。"), [])
        self.assertEqual(proofread.suspect_terms(
            "作品发表于History of Political Thought、Polity等期刊。"), [])
        self.assertEqual(proofread.suspect_terms(
            "题为Aristote Politique VII: La constitution “selon nos vœux”的特刊。"), [])
        self.assertEqual(proofread.suspect_terms(
            "原初希腊语中的 polis、politeia 和 politikos 被改编为拉丁语。"), [])
        self.assertEqual(proofread.suspect_terms(
            "历史串联于城邦、帝国与 le doux commerce 之间。"), [])

    def test_english_sentence_left_in_chinese_is_flagged(self):
        from llmtr import proofread
        self.assertTrue(proofread.ascii_prose("这是那种 launching a thousand ships 的思想。"))
        self.assertFalse(proofread.ascii_prose("发表于Journal of Political Philosophy。"))
        self.assertFalse(proofread.ascii_prose("历史串联于 le doux commerce 之间。"))
        self.assertFalse(proofread.ascii_prose("拉丁原文为 neminem posse opibus virtuti simul libertas。"))

    def test_glossary_terms_match_whole_words(self):
        from llmtr import translate
        blocks = [{"id": "b0001", "type": "p", "text": "The striking continuity of legitimacy."}]
        prompt = translate.build_prompt(blocks, {"TRI": "三", "NU": "努", "legitimacy": "合法性"})
        self.assertIn("legitimacy=合法性", prompt)
        self.assertNotIn("TRI=", prompt)
        self.assertNotIn("NU=", prompt)

    def test_emphasis_markers_become_italics(self):
        from docx import Document

        from llmtr import build
        doc = Document()
        p = doc.add_paragraph()
        build._add_marked(p, "著作包括 *Edmund Burke and the Invention*（Oxford, 2017）。")
        self.assertNotIn("*", "".join(r.text for r in p.runs))
        italics = [r.text for r in p.runs if r.italic]
        self.assertEqual(italics, ["Edmund Burke and the Invention"])

    def test_finalize_keeps_single_reference_heading(self):
        from llmtr import clean
        blocks = [
            {"page": "5", "type": "h2", "text": "References", "zh": "参考文献"},
            {"page": "5", "type": "bib", "text": "Choat, Simon. 2023. Introduction."},
            {"page": "6", "type": "p", "text": "Body.", "zh": "正文。"},
        ]
        out = clean.finalize(blocks, None)
        self.assertEqual([b["type"] for b in out].count("h1"), 0)

    def test_toc_apply_skips_existing_section_heading(self):
        from llmtr import toc
        blocks = [
            {"page": "5", "type": "h2", "text": "References", "zh": "参考文献"},
            {"page": "5", "type": "p", "text": "Entry.", "zh": "条目。"},
        ]
        out = toc.apply(blocks, [{"title": "References", "page": "5", "level": 1}],
                        lambda titles: {})
        self.assertEqual([b["type"] for b in out].count("h1"), 0)

    def test_name_candidates_drop_sentence_words(self):
        from llmtr import names
        blocks = [{"type": "p", "text": "More than this. More than that. Green (1995) and "
                                        "Green (2013) and Green 2020. More of it. More still. "
                                        + "more " * 13}]
        got = names.candidates(blocks, 150)
        self.assertIn("Green", got)
        self.assertNotIn("More", got)

    def test_name_records_mark_given_names(self):
        from llmtr import names

        class Fake:
            def chat(self, system, prompt, stop=None):
                return "John\t约翰\t名\nGreen\t格林\t姓\nElgar\t-\t\n"

        recs = names.translate_records(Fake(), ["John", "Green", "Elgar"])
        self.assertEqual(recs["John"]["kind"], names.GIVEN)
        self.assertEqual(recs["Green"]["kind"], names.SURNAME)
        self.assertNotIn("Elgar", recs)
        rows = names.build_index({"John": "约翰", "Green": "格林"},
                                 {"John": {"4"}, "Green": {"5"}}, given={"John"})
        self.assertEqual(rows, ["Green\u3000格林\u30005"])

    def test_cited_names_confirm_surnames(self):
        from llmtr import names
        blocks = [{"type": "p", "text": "See Alexander 2014 and Vincent (1997). "
                                        "John Stuart Mill wrote about liberty."}]
        got = names.cited_names(blocks)
        self.assertIn("Alexander", got)
        self.assertIn("Vincent", got)
        self.assertNotIn("John", got)

    def test_localise_leaves_citations_in_latin(self):
        from llmtr import names
        blocks = [{"type": "p", "page": "12",
                   "text": "Mark Warren (1989, 606) wrote; see Warren 1989.",
                   "zh": "马克·沃伦（Mark Warren，1989, 606）写道；参见 Warren 1989，606。"}]
        names.localise(blocks, {"Warren": "沃伦"})
        self.assertIn("（Mark Warren，1989, 606）", blocks[0]["zh"])
        self.assertIn("Warren 1989", blocks[0]["zh"])
        self.assertNotIn("Mark 沃伦", blocks[0]["zh"])

    def test_localise_keeps_book_titles_english(self):
        from llmtr import names
        blocks = [{"type": "p", "page": "1", "text": "Book.",
                   "zh": "他的著作包括 John of Paris: Beyond Royal and Papal Power（Brepols, 2015）。"}]
        names.localise(blocks, {"John": "约翰"})
        self.assertIn("John of Paris", blocks[0]["zh"])
        blocks = [{"type": "p", "page": "1", "text": "Book.",
                   "zh": "她是 Edmund Burke and the Invention 的作者。"}]
        names.localise(blocks, {"Burke": "伯克"})
        self.assertIn("Edmund Burke", blocks[0]["zh"])

    def test_localise_translates_prose_names(self):
        from llmtr import names
        blocks = [{"type": "p", "page": "12", "text": "Warren wrote.",
                   "zh": "Warren 写道。"}]
        names.localise(blocks, {"Warren": "沃伦"})
        self.assertEqual(blocks[0]["zh"], "沃伦（Warren） 写道。")

    def test_localise_can_use_the_full_source_name(self):
        from llmtr import names
        blocks = [{"type": "p", "page": "12", "text": "Mark Warren wrote.",
                   "zh": "Warren 写道。"}]
        names.localise(blocks, {"Warren": "沃伦"})
        self.assertEqual(blocks[0]["zh"], "沃伦（Mark Warren） 写道。")

    def test_localise_strong_mode_glosses_each_chapter(self):
        from llmtr import names
        blocks = [
            {"type": "h1", "page": "1", "text": "One", "zh": "第一章"},
            {"type": "p", "page": "2", "text": "Warren wrote.", "zh": "Warren 写道。"},
            {"type": "h1", "page": "9", "text": "Two", "zh": "第二章"},
            {"type": "p", "page": "10", "text": "Warren wrote again.", "zh": "Warren 又写道。"},
        ]
        names.localise(blocks, {"Warren": "沃伦"}, mode="strong")
        self.assertIn("（Warren）", blocks[1]["zh"])
        self.assertIn("（Warren）", blocks[3]["zh"])

        blocks[1]["zh"], blocks[3]["zh"] = "Warren 写道。", "Warren 又写道。"
        names.localise(blocks, {"Warren": "沃伦"}, mode="light")
        self.assertIn("（Warren）", blocks[1]["zh"])
        self.assertEqual(blocks[3]["zh"], "沃伦 又写道。")

        blocks[1]["zh"], blocks[3]["zh"] = "Warren 写道。", "Warren 又写道。"
        names.localise(blocks, {"Warren": "沃伦"}, mode="off")
        self.assertEqual(blocks[1]["zh"], "沃伦 写道。")
        self.assertEqual(blocks[3]["zh"], "沃伦 又写道。")

    def test_toc_promotes_front_matter_sections(self):
        from llmtr import build
        blocks = [
            {"type": "h2", "page": "1", "text": "About the editors", "zh": "关于编者"},
            {"type": "h2", "page": "2", "text": "Contributors", "zh": "撰稿人"},
            {"type": "h1", "page": "9", "text": "Introduction", "zh": "导论"},
        ]
        got = build.toc_heads(blocks, {})
        self.assertEqual([b["type"] for b in got], ["h1", "h1", "h1"])
        self.assertEqual(got[1]["zh"], "撰稿人")


class SourceGlossModes(unittest.TestCase):
    def test_prompt_has_three_distinct_levels(self):
        from llmtr import gloss
        off, light, strong = (gloss.system_prompt(m) for m in gloss.MODES)
        self.assertIn("不加英文原文括号", off)
        self.assertIn("首次出现加原文括号", light)
        self.assertIn("每章第一次出现都括注原文", strong)
        self.assertNotEqual(strong, light)
        self.assertNotEqual(light, off)

    def test_default_is_strong(self):
        from llmtr import gloss
        self.assertEqual(gloss.DEFAULT_MODE, "strong")
        self.assertEqual(gloss.mode_of({}), "strong")
        self.assertEqual(gloss.mode_of({"source_gloss": {"mode": "light"}}), "light")
        self.assertEqual(gloss.mode_of({"source_gloss": "bogus"}), "strong")

    def test_glossary_term_is_glossed_once(self):
        from llmtr import gloss
        cfg = {"source_gloss": {"mode": "strong"}}
        blocks = [
            {"type": "p", "text": "The empire rose.", "zh": "这个帝国兴起了。"},
            {"type": "p", "text": "The empire fell.", "zh": "这个帝国衰落了。"},
        ]
        stats = gloss.apply(blocks, {"empire": "帝国"}, cfg)
        self.assertEqual(stats["terms"], 1)
        self.assertEqual(blocks[0]["zh"], "这个帝国（empire）兴起了。")
        self.assertEqual(blocks[1]["zh"], "这个帝国衰落了。")

    def test_term_only_inside_a_gloss_waits_for_a_plain_occurrence(self):
        from llmtr import gloss
        cfg = {"source_gloss": {"mode": "light"}}
        blocks = [
            {"type": "p", "text": "The empire rose.", "zh": "（帝国）一词很含糊。"},
            {"type": "p", "text": "The empire fell.", "zh": "帝国衰落了。"},
        ]
        gloss.apply(blocks, {"empire": "帝国"}, cfg)
        self.assertEqual(blocks[0]["zh"], "（帝国）一词很含糊。")
        self.assertEqual(blocks[1]["zh"], "帝国（empire）衰落了。")

    def test_off_mode_leaves_terms_and_titles_alone(self):
        from llmtr import gloss
        cfg = {"source_gloss": {"mode": "off"}}
        blocks = [{"type": "p", "text": "The empire rose.",
                   "zh": "《帝国》（The Empire）讲这个帝国。"}]
        stats = gloss.apply(blocks, {"empire": "帝国"}, cfg)
        self.assertEqual(stats["terms"], 0)
        self.assertEqual(blocks[0]["zh"], "《帝国》讲这个帝国。")

    def test_title_gloss_is_kept_once_per_scope(self):
        from llmtr import gloss
        cfg = {"source_gloss": {"mode": "light"}}
        blocks = [
            {"type": "p", "text": "A.", "zh": "《社会契约论》（The Social Contract）是名著。"},
            {"type": "p", "text": "B.", "zh": "重读《社会契约论》（The Social Contract）很有收获。"},
        ]
        gloss.apply(blocks, {}, cfg)
        self.assertIn("（The Social Contract）", blocks[0]["zh"])
        self.assertEqual(blocks[1]["zh"], "重读《社会契约论》很有收获。")

        blocks[0]["zh"] = "《社会契约论》（The Social Contract）是名著。"
        blocks[1]["zh"] = "重读《社会契约论》（The Social Contract）很有收获。"
        cfg = {"source_gloss": {"mode": "strong"}}
        blocks.insert(1, {"type": "h1", "text": "Two", "zh": "第二章"})
        gloss.apply(blocks, {}, cfg)
        self.assertIn("（The Social Contract）", blocks[2]["zh"])

    def test_translate_blocks_passes_the_mode_prompt(self):
        from llmtr import gloss, translate
        seen = []

        class Fake:
            key = "k"

            def chat(self, system, prompt, stop=None):
                seen.append(system)
                return "[b0000] 译文"

        blocks = [{"id": "b0000", "type": "p", "text": "Hello world."}]
        translate.translate_blocks(Fake(), blocks, {}, workers=1, source_gloss="off")
        self.assertEqual(seen[0], gloss.system_prompt("off"))


class TocHeadingLevels(unittest.TestCase):
    def test_indented_toc_entries_are_sections(self):
        from unittest import mock

        from llmtr import toc
        raw = [
            {"title": "1 Crossing the Boundary", "page": "16", "chapter": True,
             "level": 1, "x0": 54.0},
            {"title": "The Tumen River Region", "page": "17", "chapter": False,
             "level": 1, "x0": 80.0},
            {"title": "2 Dynastic Geography", "page": "54", "chapter": True,
             "level": 1, "x0": 54.0},
            {"title": "I.", "page": "17", "chapter": False, "level": 1, "x0": 54.0},
        ]
        with mock.patch("llmtr.ingest._parse_contents", return_value=raw):
            got = toc.find_entries("x.pdf")
        self.assertEqual([e["level"] for e in got], [1, 2, 1])
        self.assertEqual([e["title"] for e in got],
                         ["Crossing the Boundary", "The Tumen River Region",
                          "Dynastic Geography"])

    def test_numbered_depth_sets_heading_level(self):
        from llmtr import toc
        entries = [
            {"title": "Economy", "page": "5", "level": 1},
            {"title": "Markets", "page": "5", "level": 2},
            {"title": "Prices", "page": "6", "level": 3},
        ]
        blocks = [
            {"page": "5", "type": "h1", "text": "2 Economy", "zh": "经济"},
            {"page": "5", "type": "h1", "text": "Markets", "zh": "市场"},
            {"page": "6", "type": "h2", "text": "6.3.1 Prices", "zh": "价格"},
        ]
        out = toc.apply(blocks, entries, lambda ts: {})
        got = [(b["text"], b["type"]) for b in out if b["type"][0] == "h"]
        self.assertEqual(got, [("2 Economy", "h1"), ("Markets", "h2"),
                               ("6.3.1 Prices", "h3")])

    def test_font_promoted_subsection_is_demoted(self):
        from llmtr import toc
        entries = [
            {"title": "Introduction", "page": "1", "level": 1},
            {"title": "The State and Its Rivals", "page": "2", "level": 2},
            {"title": "Economy", "page": "5", "level": 1},
            {"title": "Markets in Early Modern Europe", "page": "5", "level": 2},
        ]
        blocks = [
            {"page": "1", "type": "h1", "text": "Introduction", "zh": "导论"},
            {"page": "2", "type": "h1", "text": "The State and Its Rivals", "zh": "国家"},
            {"page": "2", "type": "p", "text": "Body text.", "zh": "正文。"},
            {"page": "5", "type": "h1", "text": "2 Economy", "zh": "经济"},
            {"page": "5", "type": "h1", "text": "Markets in Early Modern Europe", "zh": "市场"},
        ]
        out = toc.apply(blocks, entries, lambda ts: {})
        heads = [(b["text"], b["type"]) for b in out if b["type"] in ("h1", "h2")]
        self.assertEqual(heads, [("Introduction", "h1"),
                                 ("The State and Its Rivals", "h2"),
                                 ("2 Economy", "h1"),
                                 ("Markets in Early Modern Europe", "h2")])

    def test_existing_section_heading_is_not_duplicated(self):
        from llmtr import toc
        entries = [
            {"title": "Introduction", "page": "1", "level": 1},
            {"title": "The Tumen River Region", "page": "17", "level": 2},
        ]
        blocks = [
            {"page": "1", "type": "h1", "text": "Introduction", "zh": "导论"},
            {"page": "17", "type": "p", "text": "Body.", "zh": "正文。"},
            {"page": "17", "type": "h1",
             "text": "The Tumen River Region: Northeast Qing", "zh": "图们江区域"},
            {"page": "17", "type": "p", "text": "More.", "zh": "更多。"},
        ]
        out = toc.apply(blocks, entries, lambda ts: {})
        got = [(b["type"], b["text"]) for b in out
               if b["type"] in ("h1", "h2", "h3")]
        self.assertEqual(got, [("h1", "Introduction"),
                               ("h2", "The Tumen River Region: Northeast Qing")])


class GuiProgressContract(unittest.TestCase):
    def test_progress_callback_accepts_stage_index(self):
        import threading
        from types import SimpleNamespace
        from llmtr import app

        fake = SimpleNamespace(
            prog={"stage": "", "done": 0, "total": 1, "idx": 0,
                  "n": len(pipeline.STAGES), "t": 0.0},
            prog_lock=threading.Lock())
        app.App.on_progress(fake, "翻译", 3, 10, 4, 10)
        self.assertEqual(fake.prog["stage"], "翻译")
        self.assertEqual((fake.prog["done"], fake.prog["total"]), (3, 10))
        self.assertEqual((fake.prog["idx"], fake.prog["n"]), (4, 10))

    def test_render_progress_shows_stage_and_clock(self):
        import threading
        from types import SimpleNamespace
        from llmtr import app

        shown = {"bar": [], "text": []}

        bar = SimpleNamespace(configure=lambda **kw: shown["bar"].append(kw["value"]))
        var = SimpleNamespace(set=lambda v: shown["text"].append(v))
        label = SimpleNamespace(configure=lambda **kw: shown["text"].append(kw["text"]))

        fake = SimpleNamespace(
            prog={"stage": "翻译", "done": 3, "total": 10, "idx": 4,
                  "n": len(pipeline.STAGES), "t": 0.0},
            prog_lock=threading.Lock(), run_t0=0.0, stopping=False,
            bar=bar, status=var, meter=var, stage_strip=label,
            fmt_clock=app.App.fmt_clock, elapsed=lambda: 10.0)
        app.App.render_progress(fake)
        self.assertEqual(shown["bar"][0],
                         round((4 - 1 + 3 / 10) / len(pipeline.STAGES) * 100, 1))
        self.assertTrue(any("阶段 4/%d 翻译" % len(pipeline.STAGES) in t for t in shown["text"]))
        self.assertTrue(any("已用 00:10" in t for t in shown["text"]))
        self.assertTrue(any("▶4 翻译" in t for t in shown["text"]))


class ReferenceFigureGuards(unittest.TestCase):
    def test_citation_heavy_prose_is_not_a_bibliography(self):
        blocks = []
        for i in range(9):
            blocks.append({
                "page": str(260 + i), "type": "p",
                "text": "This chapter argues that contentious politics shaped state formation, as "
                        "Goldstone (1991), Wickham-Crowley (1992), and Goodwin (1994, 2001) show. "
                        "The literature from Cambridge University Press and the Journal of "
                        "Comparative Politics is extensive and the argument continues below."})
        out = ingest._detect_references(blocks)
        self.assertEqual([b["type"] for b in out], ["p"] * 9)

    def test_figure_labels_are_absorbed_but_prose_is_kept(self):
        blocks = [
            {"page": "221", "type": "note", "text": "Line separating domestic from international politics"},
            {"page": "221", "type": "note", "text": "Second Image Effects"},
            {"page": "221", "type": "p", "text": "figure 9.1. Types of domestic-international linkages"},
            {"page": "221", "type": "p", "text": "peak associations and labor unions respond to effects."},
            {"page": "220", "type": "p", "text":
                "The relationship entails two linkages (see Figure 9.1). The first are known as effects."},
        ]
        out = ingest._absorb_figure_labels(blocks)
        self.assertEqual([(b["type"], b["text"][:9]) for b in out],
                         [("p", "figure 9."), ("p", "peak asso"), ("p", "The relat")])

    def test_plate_detection_requires_low_text_volume(self):
        from llmtr import figures
        prose = {"page": "221", "type": "p", "text": "word " * 400}
        labels = [{"page": "221", "type": "note", "text": "label"} for _ in range(20)]
        self.assertNotIn(221, figures.detect(labels + [prose], dict(figures.DEFAULTS)))
        plate = [{"page": "5", "type": "note", "text": "label"} for _ in range(20)]
        self.assertIn(5, figures.detect(plate, dict(figures.DEFAULTS)))

    def test_tail_reference_pages_need_dense_entries(self):
        blocks = [{"page": str(i), "type": "p",
                   "text": "As argued, see Goldstone (1991) and Cambridge University Press."}
                  for i in range(10)]
        out = ingest._mark_tail_references(blocks)
        self.assertTrue(all(b["type"] == "p" for b in out))


class TranslationProgress(unittest.TestCase):
    def test_progress_counts_blocks_not_groups(self):
        from llmtr import translate

        blocks = [{"id": "b%04d" % i, "type": "p", "text": "word " * 5} for i in range(6)]
        seen = []

        class Fake:
            key = "k"

            def chat(self, system, prompt, stop=None):
                ids = [ln.split("]")[0][1:] for ln in prompt.splitlines() if ln.startswith("[b")]
                return "\n".join("[%s] 译文" % i for i in ids)

        translate.translate_blocks(Fake(), blocks, {}, workers=2, chunk_size=10,
                                   progress=lambda d, n: seen.append((d, n)))
        self.assertEqual(seen[-1], (6, 6))
        self.assertTrue(all(n == 6 for _d, n in seen))


class ConfigDefaults(unittest.TestCase):
    def test_shipped_config_has_no_book_specific_overrides(self):
        cfg = json.loads((ROOT / "config" / "config.json").read_text(encoding="utf-8"))
        self.assertEqual(cfg["figures"]["force"], [])
        self.assertEqual(cfg["figures"]["rotate"], {})
        self.assertEqual(cfg["figures"]["captions"], {})
        self.assertEqual(cfg["ocr"]["spread"], "auto")
        self.assertEqual(cfg["source_gloss"]["mode"], "strong")

    def test_defaults_disable_no_ocr(self):
        cfg = pipeline.load_config(None)
        self.assertEqual(cfg["ocr"]["mode"], "auto")
        self.assertEqual(cfg["figures"].get("force", []), [])
        self.assertEqual(cfg["source_gloss"]["mode"], "strong")


class CooperativeStop(unittest.TestCase):
    """Aborting must unwind quickly instead of waiting for in-flight API calls."""

    def test_parallel_stops_without_waiting_for_inflight(self):
        import threading
        import time

        from llmtr import translate

        stop = threading.Event()
        threading.Thread(target=lambda: (time.sleep(0.2), stop.set()), daemon=True).start()
        t0 = time.monotonic()
        with self.assertRaises(translate.StopRequested):
            translate._parallel([lambda: time.sleep(5)] * 4, workers=4, stop=stop)
        self.assertLess(time.monotonic() - t0, 2.0)

    def test_translate_blocks_honours_preset_stop(self):
        import threading

        from llmtr import translate

        stop = threading.Event()
        stop.set()
        blocks = [{"id": "b0000", "type": "p", "text": "Hello world."}]
        client = translate.Client("https://example.invalid/v1", "k", "m")
        with self.assertRaises(translate.StopRequested):
            translate.translate_blocks(client, blocks, {}, stop=stop)

    def test_pipeline_short_circuits_when_stopped(self):
        import threading

        pdf = ROOT.parent / "Sample12.pdf"
        if not pdf.exists():
            self.skipTest("sample PDF not available")
        stop = threading.Event()
        stop.set()
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / "out"
            res = pipeline.run(pdf, out, pipeline.load_config(None), dry_run=True,
                               stop=stop, log=lambda *_a: None)
            self.assertEqual(res.get("reason"), "stopped")
            self.assertFalse(out.exists())


def _jpeg() -> bytes:
    import io

    buf = io.BytesIO()
    Image.new("RGB", (300, 200), "white").save(buf, "JPEG")
    return buf.getvalue()


def _epub_page(*cells: str) -> str:
    rows = "".join("<tr><td>%s</td></tr>" % c for c in cells)
    return ('<?xml version="1.0" encoding="utf-8"?><html xmlns="http://www.w3.org/1999/xhtml">'
            "<head><title>x</title></head><body><table>%s</table></body></html>" % rows)


def _build_epub(root: Path) -> Path:
    import zipfile

    def page(n, *cells):
        return _epub_page('<div align="center"><font size="2" color="#FF0000">Page %s</font></div>' % n,
                          *cells)

    prose = ("The railway carried grain, timber and coal across a plain that had never "
             "known a wheeled vehicle before. ") * 3
    files = {
        "mimetype": "application/epub+zip",
        "META-INF/container.xml":
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles>'
            '<rootfile full-path="OEBPS/content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>',
        "OEBPS/content.opf":
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="2.0" '
            'unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
            "<dc:title>A Test Book : With a Subtitle</dc:title>"
            "<dc:creator>Doe, Jane</dc:creator><dc:publisher>Test Press</dc:publisher>"
            "</metadata><manifest>"
            + "".join('<item id="p%s" href="page_%s.xhtml" media-type="application/xhtml+xml"/>'
                      % (i, i) for i in ("i", 1, 2, 3, 4, 5, 6))
            + '<item id="cover" href="cover.xhtml" media-type="application/xhtml+xml"/>'
            "</manifest><spine>"
            + "".join('<itemref idref="p%s"/>' % i for i in ("i", 1, 2, 3, 4, 5, 6))
            + "</spine></package>",
        "OEBPS/cover.xhtml":
            "<html><body><table>"
            "<tr><td>Title</td><td>:</td><td>A Test Book</td></tr>"
            "<tr><td>Publication Date</td><td>:</td><td>1999</td></tr>"
            "</table></body></html>",
        "OEBPS/page_i.xhtml": page(
            "i", '<font size="5">CONTENTS</font>',
            '<font size="3">Introduction ........ 1</font>'),
        "OEBPS/page_1.xhtml": page(
            1, '<font size="5">INTRODUCTION</font>',
            '<font size="3">%s a plain that would' % prose),
        "OEBPS/page_2.xhtml": page(
            2, '<font size="3">soon be crossed by the first locomotive.</font>'),
        "OEBPS/page_3.xhtml": page(
            3, '<img src="images/fig.jpg" width="300" height="200"/>',
            '<img src="images/spacer.gif" width="0" height="1"/>',
            '<font size="4">PART ONE</font>',
            '<font size="5">1 BEGINNINGS</font>',
            '<font size="3">%s</font>' % prose,
            '<font size="2">1. A chapter note about the railway.</font>',
            '<font size="2">2. Another note, on grain tariffs.</font>'),
        "OEBPS/page_4.xhtml": page(
            4, '<font size="4">REFERENCE NOTES</font>',
            '<font size="3">Chapter I: Beginnings</font>',
            '<font size="3">The first reference note, longer than a heading line.</font>'),
        "OEBPS/page_5.xhtml": page(
            5, '<font size="4">BIBLIOGRAPHY</font>',
            '<font size="3">Smith, John. A History of Wheels. London: Test Press, 1990.</font>',
            '<font size="3">Jones, Ada. Rails and Roads. Oxford: Test Press, 1995.</font>'),
        "OEBPS/page_6.xhtml": page(
            6, '<font size="4">INDEX</font>',
            '<font size="3">Smith, John, 1, 3</font>',
            '<font size="3">Rails, 2</font>'),
        "OEBPS/images/fig.jpg": _jpeg(),
    }
    path = root / "test_book.epub"
    with zipfile.ZipFile(path, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)
    return path


class EpubIngest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from llmtr import epub

        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.tmp.name)
        cls.path = _build_epub(cls.root)
        cls.book = epub.extract(cls.path, cls.root / "assets", log=lambda *_a: None)
        cls.blocks = cls.book["blocks"]
        cls.kinds = Counter(b["type"] for b in cls.blocks)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_metadata_from_opf_and_cover(self):
        meta = self.book["meta"]
        self.assertEqual(meta["title"], "A Test Book")
        self.assertEqual(meta["subtitle"], "With a Subtitle")
        self.assertEqual(meta["author"], "Jane Doe")
        self.assertEqual(meta["publisher"], "Test Press")
        self.assertEqual(meta["year"], "1999")

    def test_printed_contents_page_is_skipped(self):
        self.assertEqual(self.book["labels"], ["i", "1", "2", "3", "4", "5", "6"])
        self.assertFalse(any("CONTENTS" in b["text"].upper() for b in self.blocks))
        self.assertFalse(any("........" in b["text"] for b in self.blocks))

    def test_heading_levels_follow_the_book(self):
        h1 = [b["text"] for b in self.blocks if b["type"] == "h1"]
        self.assertEqual(h1, ["INTRODUCTION", "PART ONE", "REFERENCE NOTES",
                              "BIBLIOGRAPHY", "INDEX"])
        self.assertEqual([b["text"] for b in self.blocks if b["type"] == "h2"],
                         ["1 BEGINNINGS", "NOTES"])

    def test_notes_bibliography_and_index(self):
        self.assertGreaterEqual(self.kinds["note"], 3)
        self.assertGreaterEqual(self.kinds["bib"], 4)
        self.assertIn("Chapter I: Beginnings", [b["text"] for b in self.blocks
                                               if b["type"] == "h3"])
        note_head = [b for b in self.blocks if b["type"] == "h2" and b["text"] == "NOTES"]
        self.assertTrue(note_head and note_head[0].get("endnote"))
        ref = next(b for b in self.blocks if b["text"] == "REFERENCE NOTES")
        self.assertTrue(ref.get("endnote"))

    def test_cross_page_sentence_is_rejoined(self):
        intro = next(b["text"] for b in self.blocks
                     if b["type"] == "p" and b["text"].startswith("The railway"))
        self.assertIn("would soon be crossed", intro)
        self.assertEqual(next(b for b in self.blocks if b["text"].startswith("The railway"))["page_end"], "2")

    def test_only_real_images_become_figures(self):
        figs = [b for b in self.blocks if b["type"] == "figure"]
        self.assertEqual(len(figs), 1)
        self.assertTrue(Path(figs[0]["image"]).exists())
        self.assertEqual(Path(figs[0]["image"]).suffix, ".jpg")


class EpubPipeline(unittest.TestCase):
    def test_dry_run_builds_docx_without_inventing_a_name_index(self):
        from unittest import mock

        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            path = _build_epub(root)
            out = root / "out"
            with mock.patch.object(pipeline, "paginate", return_value={}):
                res = pipeline.run(path, out, pipeline.load_config(None), dry_run=True,
                                   log=lambda *_a: None)
            self.assertTrue(res["ok"], res)
            self.assertTrue(res["docx"].exists())
            book = json.loads((out / "book.json").read_text(encoding="utf-8"))
            h1 = [b["zh"] for b in book["blocks"] if b["type"] == "h1"]
            self.assertNotIn("人名索引", h1)
            self.assertTrue(any("索引" in t or "INDEX" in t for t in h1))


class DocxIngest(unittest.TestCase):
    """Word styles carry the outline: read it instead of guessing font sizes."""

    @classmethod
    def setUpClass(cls):
        from docx import Document

        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.tmp.name)
        cls.path = cls.root / "book.docx"
        doc = Document()
        doc.core_properties.title = "Test Book: A Subtitle"
        doc.core_properties.author = "Jane Doe"
        from datetime import datetime
        doc.core_properties.created = datetime(1999, 5, 4)
        doc.add_paragraph("Test Book: A Subtitle", style="Title")
        doc.add_paragraph("A Subtitle", style="Subtitle")
        doc.add_paragraph("Jane Doe")
        doc.add_paragraph("Chapter One", style="Heading 1")
        doc.add_paragraph("Body text that is not a heading.")
        doc.add_paragraph("Background", style="Heading 2")
        doc.add_paragraph("Details", style="Heading 3")
        doc.add_paragraph("Contents", style="Heading 1")
        doc.save(str(cls.path))
        cls._with_company(cls.path, "Test Press")

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    @staticmethod
    def _with_company(path: Path, company: str) -> None:
        import zipfile

        tmp = path.with_suffix(".tmp")
        xml = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
               'extended-properties"><Company>%s</Company></Properties>' % company)
        with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "docProps/app.xml":
                    data = xml.encode("utf-8")
                zout.writestr(item, data)
        tmp.replace(path)

    def test_headings_come_from_styles(self):
        from llmtr import docxin

        heads = docxin.headings(self.path)
        self.assertEqual([(h["level"], h["text"]) for h in heads],
                         [(1, "Chapter One"), (2, "Background"), (3, "Details")])

    def test_core_metadata(self):
        from llmtr import docxin

        meta = docxin.core_meta(self.path)
        self.assertEqual(meta["title"], "Test Book")
        self.assertEqual(meta["subtitle"], "A Subtitle")
        self.assertEqual(meta["author"], "Jane Doe")
        self.assertEqual(meta["year"], "1999")
        self.assertEqual(meta["publisher"], "Test Press")

    def test_front_matter_is_dropped_and_ids_reindexed(self):
        from llmtr import docxin

        blocks = [{"id": "b0000", "page": "1", "type": "h2", "text": "Test Book"},
                  {"id": "b0001", "page": "1", "type": "p", "text": "A Subtitle"},
                  {"id": "b0002", "page": "1", "type": "byline", "text": "Jane Doe"},
                  {"id": "b0003", "page": "1", "type": "h1", "text": "Chapter One"}]
        meta = {"title": "Test Book", "subtitle": "A Subtitle", "author": "Jane Doe"}
        self.assertEqual(docxin.drop_front_matter(blocks, meta), 3)
        self.assertEqual([b["text"] for b in blocks], ["Chapter One"])
        self.assertEqual(blocks[0]["id"], "b0000")

    def test_stamp_outline_including_merged_heading(self):
        from llmtr import docxin

        blocks = [
            {"page": "1", "type": "h2", "text": "CHAPTER ONE"},
            {"page": "1", "type": "p",
             "text": "Recruitment Criteria Recruitment combined political loyalty "
                     "with academic performance for all incoming students."},
            {"page": "1", "type": "p", "text": "The Party Organization"},
            {"page": "1", "type": "figure", "text": ""},
            {"page": "1", "type": "p", "text": "Details"},
        ]
        heads = [{"level": 1, "text": "Chapter One"},
                 {"level": 2, "text": "Recruitment Criteria"},
                 {"level": 2, "text": "The Party Organization"},
                 {"level": 3, "text": "Details"}]
        self.assertEqual(docxin.apply_headings(blocks, heads), 4)
        self.assertEqual([b["type"] for b in blocks], ["h1", "h2", "p", "h2", "figure", "h3"])
        self.assertEqual(blocks[1]["text"], "Recruitment Criteria")
        self.assertTrue(blocks[2]["text"].startswith("Recruitment combined"))
        self.assertEqual(blocks[4]["type"], "figure")

    def test_short_prefix_remainder_is_not_split(self):
        from llmtr import docxin

        blocks = [{"page": "1", "type": "p", "text": "Summary Appendix"}]
        self.assertEqual(docxin.apply_headings(blocks, [{"level": 3, "text": "Summary"}]), 0)
        self.assertEqual(blocks[0]["type"], "p")

    def test_convert_falls_back_to_office_com(self):
        from unittest import mock

        from llmtr import docxin

        src = self.root / "convert.docx"
        src.write_bytes(b"x")
        out = self.root / "convert.pdf"

        def fake(src_, out_, stop=None):
            Path(out_).write_bytes(b"%PDF-1.4")
            return True

        with mock.patch.object(docxin, "_word_com", side_effect=fake):
            self.assertTrue(docxin.convert(src, out, node=None, cli=None, soffice=None,
                                           log=lambda *_a: None))
        self.assertTrue(out.exists())

    def test_convert_uses_soffice_output(self):
        from unittest import mock

        from llmtr import docxin

        src = self.root / "soffice.docx"
        src.write_bytes(b"x")
        out = self.root / "soffice.pdf"

        def fake_run(cmd, stop=None, timeout=None):
            Path(cmd[-1]).with_suffix(".pdf").write_bytes(b"%PDF-1.4")
            return True

        with mock.patch.object(docxin, "_run", side_effect=fake_run):
            self.assertTrue(docxin.convert(src, out, node=None, cli=None, soffice="soffice",
                                           log=lambda *_a: None))
        self.assertTrue(out.exists())


class DocxPipeline(unittest.TestCase):
    def test_conversion_failure_is_reported(self):
        from unittest import mock

        from llmtr import docxin

        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            src = root / "book.docx"
            src.write_bytes(b"x")
            with mock.patch.object(docxin, "convert", return_value=False):
                res = pipeline.run(src, root / "out", pipeline.load_config(None),
                                   dry_run=True, log=lambda *_a: None)
            self.assertFalse(res["ok"])
            self.assertEqual(res["reason"], "docx-convert")


class HeadingFragments(unittest.TestCase):
    def test_page_break_title_fragment_is_dropped(self):
        blocks = [
            {"page": "162", "type": "h1", "text": "FROM COALITION TO CONFRONTATION (1962-1965)"},
            {"page": "162", "type": "h1", "text": "FROM COALITION TO"},
            {"page": "226", "type": "h1", "text": "TAKE-OVER (1966-1967)"},
        ]
        out = ingest._drop_heading_fragments([dict(b) for b in blocks])
        self.assertEqual([b["text"] for b in out],
                         ["FROM COALITION TO CONFRONTATION (1962-1965)",
                          "TAKE-OVER (1966-1967)"])

    def test_part_heading_kept_before_its_expansion(self):
        blocks = [
            {"page": "1", "type": "h1", "text": "PART ONE"},
            {"page": "1", "type": "h1", "text": "PART ONE: THE ARMY"},
        ]
        out = ingest._drop_heading_fragments([dict(b) for b in blocks])
        self.assertEqual(len(out), 2)


class ClientProtocol(unittest.TestCase):
    """Any OpenAI-compatible endpoint must work: URL, auth, model, usage, cost."""

    def test_request_matches_openai_compatible_api(self):
        from unittest import mock

        from llmtr import translate

        seen = {}

        class Resp:
            def read(self):
                return json.dumps({
                    "choices": [{"message": {"content": "你好"}}],
                    "usage": {"prompt_tokens": 1000, "completion_tokens": 500},
                }).encode("utf-8")

            def __enter__(self):
                return self

            def __exit__(self, *_a):
                return False

        def fake_urlopen(req, timeout=None):
            seen["url"] = req.full_url
            seen["auth"] = req.get_header("Authorization")
            seen["extra"] = req.get_header("X-extra")
            seen["body"] = json.loads(req.data.decode("utf-8"))
            return Resp()

        client = translate.Client("https://dashscope.aliyuncs.com/compatible-mode/v1",
                                  "sk-test", "qwen-plus", price_in=0.8, price_out=2.0,
                                  headers={"X-Extra": "1"})
        with mock.patch.object(translate.urllib.request, "urlopen", fake_urlopen):
            self.assertEqual(client.chat("sys", "user"), "你好")
        self.assertEqual(seen["url"],
                         "https://dashscope.aliyuncs.com/compatible-mode/v1/chat/completions")
        self.assertEqual(seen["auth"], "Bearer sk-test")
        self.assertEqual(seen["extra"], "1")
        self.assertEqual(seen["body"]["model"], "qwen-plus")
        self.assertEqual(client.usage["prompt"], 1000)
        self.assertAlmostEqual(client.cost(), 0.0018, places=9)

    def test_cache_errors_never_fail_a_run(self):
        from unittest import mock

        from llmtr import translate

        class Resp:
            def read(self):
                return json.dumps({"choices": [{"message": {"content": "ok"}}]}).encode("utf-8")

            def __enter__(self):
                return self

            def __exit__(self, *_a):
                return False

        with tempfile.TemporaryDirectory() as d:
            blocker = Path(d) / "cache"
            blocker.write_text("not a directory", encoding="utf-8")
            client = translate.Client("https://example.com/v1", "k")
            client.cache_dir = blocker
            self.assertIsNone(client.get_block("abc"))
            client.put_block("abc", "你好")
            with mock.patch.object(translate.urllib.request, "urlopen",
                                   lambda req, timeout=None: Resp()):
                self.assertEqual(client.chat("s", "u"), "ok")


class TableAsImage(unittest.TestCase):
    """Tables keep the book's original layout: cropped to PNG, never rebuilt."""

    def test_ruled_table_becomes_an_image(self):
        from llmtr import tables

        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            pdf = root / "t.pdf"
            doc = pymupdf.open()
            page = doc.new_page(width=400, height=600)
            rows = [["Region", "1990", "2000"],
                    ["North", "12", "18"],
                    ["South", "9", "15"],
                    ["East", "7", "11"]]
            x0, y0, cw, rh = 60, 100, 100, 24
            for r in range(len(rows) + 1):
                page.draw_line((x0, y0 + r * rh), (x0 + cw * 3, y0 + r * rh), width=0.8)
            for c in range(4):
                page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + len(rows) * rh), width=0.8)
            for r, row in enumerate(rows):
                for c, txt in enumerate(row):
                    page.insert_text((x0 + c * cw + 4, y0 + r * rh + 16), txt, fontsize=9)
            page.insert_text((60, 300), "Body prose continues below the table.", fontsize=10)
            doc.save(str(pdf))

            found = tables.candidates(pdf, ["1"], dict(tables.DEFAULTS))
            self.assertEqual(len(found), 1)
            blocks = [{"page": "1", "type": "p", "text": "Region 1990 2000 North 12 18",
                       "zh": "地区 1990 2000 北 12 18"}]
            blocks, done = tables.replace(blocks, found, pdf, ["1"], root / "assets",
                                          dict(tables.DEFAULTS), client=None)
            self.assertEqual(done, [1])
            tabs = [b for b in blocks if b["type"] == "table"]
            self.assertEqual(len(tabs), 1)
            self.assertTrue(Path(tabs[0]["image"]).exists())
            self.assertTrue(tabs[0]["caption"].startswith("表"))


class ProofLevel(unittest.TestCase):
    def test_multi_note_block_splits_on_markers(self):
        from llmtr import build

        parts = build._note_parts(
            "\u27e618\u27e7 Lattimore, Inner Asian Frontiers, p. 105. "
            "\u27e619\u27e7 See also Rawski, The Last Emperors.", "")
        self.assertEqual([p[0] for p in parts], ["18", "19"])
        self.assertTrue(parts[0][1].startswith("Lattimore"))
        self.assertTrue(parts[1][1].startswith("See also"))

    def test_note_lookup_does_not_cross_numeric_pages(self):
        from llmtr import build

        recs = [{"no": "2", "text": "x", "page": "16", "used": False}]
        self.assertEqual(build._note_lookup(recs, "2", "16"), (recs[0], "exact"))
        self.assertEqual(build._note_lookup(recs, "2", "17"), (recs[0], "near"))
        self.assertEqual(build._note_lookup(recs, "2", "xvii"), (None, None))

    def test_explicit_figure_number_is_preserved(self):
        from llmtr import build

        blocks = [{"type": "figure", "image": "x.png", "page": "221",
                   "caption": "figure 9.1. Types of linkages"}]
        refs, dups = build._assign_figure_numbers(blocks)
        self.assertEqual(blocks[0]["caption"], "图 9.1\u3000Types of linkages")
        self.assertEqual(refs["fig:9.1"], blocks[0]["anchor"])
        self.assertEqual(dups, [])

    def test_proof_report_is_written(self):
        from llmtr import proofread

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            book = {"meta": {"title": "T"}, "blocks": [
                {"id": "b1", "page": "1", "type": "p", "text": "Hello", "zh": "你好"}]}
            stats = {"notes_total": 2, "notes_anchored": 2, "notes_exact": 2,
                     "figures": 1, "tables": 0, "markers_total": 2}
            md = proofread.write_report(out, book, stats=stats, preset="print",
                                        log=lambda *a: None)
            self.assertTrue(md.exists())
            text = md.read_text(encoding="utf-8")
            self.assertIn("校样报告", text)
            self.assertIn("按注号锚定 | 2", text)
            self.assertTrue((out / "proof_report.csv").exists())

    def test_metadata_skips_affiliation_line(self):
        from llmtr import meta

        title = meta._title_line([
            "Making Borders in Modern",
            "East Asia",
            "The Tumen River Demarcation, 1881\u20131919",
            "Nianshen Song",
            "University of Maryland, Baltimore County",
        ])
        self.assertEqual(title, "Making Borders in Modern East Asia")

    def test_front_matter_notes_are_kept(self):
        from llmtr import clean

        blocks = [
            {"page": "x", "type": "p", "text": "Preface " * 60, "zh": "前言" * 60},
            {"page": "x", "type": "note", "text": "1 A note in the preface.", "zh": "1 前言中的注。"},
            {"page": "1", "type": "p", "text": "Body " * 60, "zh": "正文" * 60},
        ]
        out = clean.run(blocks)
        self.assertTrue(any(b["type"] == "note" for b in out))

    def test_year_is_not_taken_from_title_range(self):
        from llmtr import meta

        text = ("Making Borders in Modern East Asia\n"
                "The Tumen River Demarcation, 1881\u20131919\n"
                "Nianshen Song\n"
                "© Nianshen Song 2018\n"
                "First published 2018\n")
        self.assertEqual(meta._heuristic(text)["year"], "2018")

    def test_page_tags_are_opt_in(self):
        import json as _json

        from docx import Document

        from llmtr import build

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "t.docx"
            book = {"meta": {"title": "T"}, "blocks": [
                {"id": "b1", "type": "p", "page": "7", "text": "Body", "zh": "正文"}]}
            build.build_doc(_json.loads(_json.dumps(book)), "print", out,
                            {"title": "T"})
            plain = "\n".join(p.text for p in Document(str(out)).paragraphs)
            self.assertNotIn("〔p.7〕", plain)
            build.build_doc(_json.loads(_json.dumps(book)), "print", out,
                            {"title": "T", "page_tags": True})
            tagged = "\n".join(p.text for p in Document(str(out)).paragraphs)
            self.assertIn("〔p.7〕", tagged)

    def test_default_preset_is_print(self):
        from llmtr import pipeline

        self.assertEqual(pipeline.DEFAULTS["preset"], "print")
        self.assertFalse(pipeline.DEFAULTS["page_tags"])
        self.assertTrue(pipeline.DEFAULTS["odd_even_header"])


class OcrQuality(unittest.TestCase):
    def test_glued_token_detection(self):
        self.assertTrue(ingest._glued_token("coercedJordantotowtheline"))
        self.assertTrue(ingest._glued_token("thebeginningof"))
        self.assertFalse(ingest._glued_token("Furthermore"))
        self.assertFalse(ingest._glued_token("MacArthur"))
        self.assertFalse(ingest._glued_token("internationalization"))
        self.assertFalse(ingest._glued_token("Jordan"))

    def test_suspect_pages_skips_clean_and_excluded(self):
        glued = "The court coercedJordantotowtheline and Lebanon's " * 20
        clean = "The court coerced Jordan to tow the line. " * 20
        self.assertEqual(
            ingest.suspect_pages([clean, glued, "short", glued], exclude={3}),
            [1])
        self.assertEqual(ingest.suspect_pages([glued, clean], max_pages=1), [0])

    def test_text_quality_flags_widespread_glue(self):
        glued = "coercedJordantotowtheline " * 80
        clean = "the court ruled that the state could not " * 80
        bad = ingest.text_quality([glued] * 4)
        self.assertFalse(bad["ok"])
        self.assertEqual(bad["anomaly"], 1.0)
        good = ingest.text_quality([clean] * 4)
        self.assertTrue(good["ok"])
        self.assertEqual(good["anomaly"], 0.0)

    def test_parse_review_verdicts(self):
        self.assertEqual(ocr.parse_review("[OK]"), ("ok", None, []))
        verdict, label, items = ocr.parse_review(
            "[FIX]\n[PAGE] 12\n[H1] The Chapter\n[P] First line.")
        self.assertEqual(verdict, "fix")
        self.assertEqual(label, "12")
        self.assertEqual(items, [("h1", "The Chapter"), ("p", "First line.")])
        self.assertEqual(ocr.parse_review("looks fine to me")[0], "unclear")
        self.assertEqual(ocr.parse_review("[OK]\n[P] stray")[0], "fix")

    def test_vision_h1_survives_unnumbered(self):
        blocks = ingest.build_ir(["Intro\nbody text"], ["1"],
                                 paras=[[("Introduction", 10), ("body text", 0)]])
        self.assertEqual([b["type"] for b in blocks], ["h1", "p"])

    def test_review_pages_wires_verdicts(self):
        from llmtr import vision

        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = Path(tmp) / "a.pdf"
            doc = pymupdf.open()
            doc.new_page(width=400, height=600)
            doc.new_page(width=400, height=600)
            doc.save(str(pdf_path))
            doc.close()
            original = vision.review_layout
            calls = []

            def fake(client, image, ocr_text, max_side=1600, stop=None, strict=False):
                calls.append(Path(image).name)
                if ocr_text == "old one":
                    return "[OK]"
                return "[FIX]\n[PAGE] 2\n[H1] Title\n[P] New text."

            vision.review_layout = fake
            try:
                res = ocr.review_pages(pdf_path, None, {"tmp_dir": tmp},
                                       ["old one", "old two"],
                                       log=lambda *a: None, only=[0, 1])
            finally:
                vision.review_layout = original
        self.assertEqual([r["verdict"] for r in res], ["ok", "fix"])
        self.assertEqual(res[1]["label"], "2")
        self.assertEqual(res[1]["items"], [("h1", "Title"), ("p", "New text.")])

    def test_review_retries_unclear_page(self):
        from llmtr import vision

        with tempfile.TemporaryDirectory() as tmp:
            pdf_path = Path(tmp) / "a.pdf"
            doc = pymupdf.open()
            doc.new_page(width=400, height=600)
            doc.save(str(pdf_path))
            doc.close()
            original = vision.review_layout
            seen = []

            def fake(client, image, ocr_text, max_side=1600, stop=None, strict=False):
                seen.append(strict)
                if strict:
                    return "[FIX]\n[PAGE] 1\n[P] Rewritten."
                return "OCR 质量很差，需要重写"

            vision.review_layout = fake
            try:
                res = ocr.review_pages(pdf_path, None, {"tmp_dir": tmp,
                                                        "judge_retries": 1},
                                       ["bad text"], log=lambda *a: None, only=[0])
            finally:
                vision.review_layout = original
        self.assertEqual(seen, [False, True])
        self.assertEqual(res[0]["verdict"], "fix")
        self.assertEqual(res[0]["items"], [("p", "Rewritten.")])

    def test_report_lists_ocr_judgement(self):
        from llmtr import proofread

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp)
            book = {"meta": {"title": "T"}, "blocks": [
                {"id": "b1", "page": "1", "type": "p", "text": "Hello", "zh": "你好"}]}
            stats = {"ocr": {"vision": False, "severe": [4], "suspect": [13, 20],
                             "fixed": [4, 20], "kept": [13], "unclear": [],
                             "skipped": [7], "judged": 2}}
            proofread.write_report(out, book, stats=stats, preset="print",
                                   log=lambda *a: None)
            text = (out / "proof_report.md").read_text(encoding="utf-8")
            self.assertIn("OCR 质量", text)
            self.assertIn("判读后保留 | 1", text)
            self.assertIn("已按图像重排：p.4、p.20", text)
            self.assertIn("OCR 判读后重排 2 页待抽查", text)
            rows = (out / "proof_report.csv").read_text(encoding="utf-8-sig")
            self.assertIn("ocr,p.20", rows)
            self.assertTrue((out / "proof_report.csv").exists())


class VisionCache(unittest.TestCase):
    """A broken cache must not lose a paid vision answer."""

    class Resp:
        def __init__(self, text):
            self.text = text

        def read(self):
            return json.dumps({"choices": [{"message": {"content": self.text}}],
                               "usage": {"prompt_tokens": 10, "completion_tokens": 2}}
                              ).encode("utf-8")

        def __enter__(self):
            return self

        def __exit__(self, *_a):
            return False

    def test_unwritable_cache_still_returns_text(self):
        from unittest import mock

        from PIL import Image

        from llmtr import vision
        from llmtr.translate import Client

        with tempfile.TemporaryDirectory() as d:
            root = Path(d)
            img = root / "p.png"
            Image.new("RGB", (16, 16), "white").save(img)
            blocker = root / "cache"
            blocker.write_text("not a directory", encoding="utf-8")
            client = Client("https://example.com/v1", "k")
            client.cache_dir = blocker
            with mock.patch.object(vision.urllib.request, "urlopen",
                                   lambda req, timeout=None: self.Resp("[OK]")):
                self.assertEqual(vision.ask(client, img, "judge"), "[OK]")

    def test_empty_answer_is_retried(self):
        from unittest import mock

        from PIL import Image

        from llmtr import vision
        from llmtr.translate import Client

        answers = [self.Resp(""), self.Resp("[OK]")]
        calls = []

        def urlopen(req, timeout=None):
            calls.append(1)
            return answers[len(calls) - 1]

        with tempfile.TemporaryDirectory() as d:
            img = Path(d) / "p.png"
            Image.new("RGB", (16, 16), "white").save(img)
            client = Client("https://example.com/v1", "k")
            with mock.patch.object(vision.urllib.request, "urlopen", urlopen), \
                    mock.patch("time.sleep", lambda _s: None):
                self.assertEqual(vision.ask(client, img, "judge"), "[OK]")
        self.assertEqual(len(calls), 2)


class ConfigPathResolution(unittest.TestCase):
    def test_paths_are_relative_to_app_root(self):
        from llmtr import app

        cfg = {"paths": {"glossary": "config/glossary.tsv", "names": "config/names.tsv"}}
        app.resolve_paths(cfg, ROOT)
        self.assertEqual(Path(cfg["paths"]["glossary"]), ROOT / "config" / "glossary.tsv")
        self.assertTrue(Path(cfg["paths"]["glossary"]).exists())
        self.assertTrue(Path(cfg["paths"]["names"]).exists())


class TranslationReliability(unittest.TestCase):
    """Long books must not lose paragraphs to output limits or stray batches."""

    @staticmethod
    def _ids(prompt):
        return [ln.split("]")[0][1:] for ln in prompt.splitlines() if ln.startswith("[b")]

    def test_truncated_tail_is_regenerated(self):
        from llmtr import translate

        blocks = [{"id": "b%04d" % i, "type": "p", "text": "alpha beta gamma %d." % i}
                  for i in range(1, 4)]
        calls = []

        class Fake:
            key = "k"

            def chat_meta(self, system, prompt, stop=None):
                calls.append(prompt)
                ids = TranslationReliability._ids(prompt)
                if len(calls) == 1:
                    body = "\n".join("[%s] 译文%s" % (i, i) for i in ids[:-1])
                    return body + "\n[%s] 未完" % ids[-1], "length"
                return "\n".join("[%s] 译文%s" % (i, i) for i in ids), "stop"

        translate.translate_blocks(Fake(), blocks, {}, workers=1, chunk_size=100)
        self.assertTrue(all(b.get("zh") for b in blocks))
        self.assertTrue(any("重试" in p for p in calls))

    def test_chunking_respects_character_budget(self):
        from llmtr import translate

        blocks = [{"id": "b%04d" % i, "type": "p",
                   "text": "x" * 2990 + "marker %04d" % i} for i in range(3)]
        calls = []

        class Fake:
            key = "k"

            def chat(self, system, prompt, stop=None):
                calls.append(prompt)
                return "\n".join("[%s] 译文" % i for i in TranslationReliability._ids(prompt))

        translate.translate_blocks(Fake(), blocks, {}, workers=1,
                                   chunk_size=100000, chunk_chars=5000)
        self.assertEqual(len(calls), 3)

    def test_unresolved_batch_is_rescued_block_by_block(self):
        from llmtr import translate

        blocks = [{"id": "b%04d" % i, "type": "p", "text": "sentence %d here." % i}
                  for i in range(1, 4)]

        class Fake:
            key = "k"

            def chat_meta(self, system, prompt, stop=None):
                ids = TranslationReliability._ids(prompt)
                if len(ids) > 1:
                    return "", "stop"
                return "[%s] 完整译文" % ids[0], "stop"

        translate.translate_blocks(Fake(), blocks, {}, workers=1, chunk_size=100)
        self.assertTrue(all(b.get("zh") == "完整译文" for b in blocks))


class PolishGuards(unittest.TestCase):
    def test_facts_and_citations_survive_any_edit(self):
        from llmtr import translate

        old = "在 1950 年，他把这一制度看做一个整体，进行了分析⟦3⟧，参见 Smith（1980）。"
        good = "1950 年，他将这一制度视为一个整体加以分析⟦3⟧，参见 Smith（1980）。"
        self.assertTrue(translate._polish_ok(old, good))
        self.assertFalse(translate._polish_ok(old, good.replace("1950", "1960")))
        self.assertFalse(translate._polish_ok(old, good.replace("⟦3⟧", "")))
        self.assertFalse(translate._polish_ok(old, good.replace("Smith", "史密斯")))
        self.assertFalse(translate._polish_ok(old, "太短。"))

    def test_polish_applies_guarded_rewrite_only(self):
        from llmtr import translate

        old = "在 1950 年，他针对这一问题进行了系统的分析，并且提出了一个关于制度变迁的重要观点。"
        good = "1950 年，他系统分析了这一问题，并就制度变迁提出了一个重要观点，影响深远。"
        blocks = [{"id": "b0001", "type": "p", "page": "1",
                   "text": "A long paragraph about 1950 and institutional change. " * 8,
                   "zh": old}]

        class Fake:
            key = "k"

            def __init__(self, reply):
                self.reply = reply

            def chat(self, system, prompt, stop=None):
                return "[b0001] " + self.reply

        n = translate.polish_blocks(Fake(good), blocks, workers=1)
        self.assertEqual(n, 1)
        self.assertEqual(blocks[0]["zh"], good)

        blocks[0]["zh"] = old
        n = translate.polish_blocks(Fake(good.replace("1950", "1960")), blocks, workers=1)
        self.assertEqual(n, 0)
        self.assertEqual(blocks[0]["zh"], old)


class TocPageRecovery(unittest.TestCase):
    @staticmethod
    def _line(*words):
        return [(x, y, x + 10, y + 10, t, 0, 0, 0) for x, y, t in words]

    def test_orphan_folio_and_chapter_number_are_reattached(self):
        from llmtr.ingest import _merge_orphan_lines

        lines = [
            self._line((70, 100, "Crossing"), (110, 100, "the"), (150, 100, "Boundary")),
            self._line((54, 106, "1.")),
            self._line((400, 114, "16")),
            self._line((70, 160, "The"), (140, 160, "Tumen"), (200, 160, "Region")),
            self._line((54, 166, "2.")),
            self._line((400, 168, "54")),
        ]
        got = [" ".join(w[4] for w in row) for row in _merge_orphan_lines(lines)]
        self.assertIn("1. Crossing the Boundary 16", got)
        self.assertIn("2. The Tumen Region 54", got)

    def test_columns_sharing_a_baseline_keep_the_contents_row(self):
        from llmtr.ingest import _merge_orphan_lines

        lines = [
            self._line((70, 100, "Conclusion"), (400, 100, "261"),
                       (600, 100, "When"), (660, 100, "I"), (690, 100, "first"),
                       (740, 100, "visited"), (820, 100, "Tsinghua")),
        ]
        got = [" ".join(w[4] for w in row) for row in _merge_orphan_lines(lines)]
        self.assertEqual(got[0], "Conclusion 261")

    def test_numbered_chapter_margin_sets_section_indent(self):
        from llmtr import toc

        raw = [
            {"title": "IV. The State", "page": "5", "level": 1, "x0": 54.0, "chapter": True},
            {"title": "Markets", "page": "6", "level": 1, "x0": 80.0, "chapter": False},
            {"title": "Prices", "page": "6", "level": 1, "x0": 54.0, "chapter": True},
        ]
        got = toc._clean_entries(raw)
        self.assertEqual([e["title"] for e in got], ["The State", "Markets", "Prices"])
        self.assertEqual([e["level"] for e in got], [1, 2, 1])

    def test_garbled_folio_sequence_is_rejected(self):
        from llmtr import toc

        bad = [{"title": "T%d" % i, "page": str(200 - i * 10), "level": 1}
               for i in range(5)]
        self.assertFalse(toc._entries_ok(bad))
        good = [{"title": "T%d" % i, "page": str(10 + i * 10), "level": 1}
                for i in range(5)]
        self.assertTrue(toc._entries_ok(good))

    def test_page_less_part_promotes_matching_body_heading(self):
        from llmtr import toc

        blocks = [
            {"page": "6", "type": "h2", "text": "PART II HISTORIOGRAPHIES", "zh": "第二部分"},
            {"page": "6", "type": "p", "text": "Body text.", "zh": "正文。"},
        ]
        out = toc.apply(blocks, [{"title": "PART II HISTORIOGRAPHIES",
                                  "page": None, "level": 1}], lambda ts: {})
        self.assertEqual(out[0]["type"], "h1")


class CompletenessChecks(unittest.TestCase):
    @staticmethod
    def _book(src, zh, page="12", typ="p"):
        return {"blocks": [{"id": "b0001", "page": page, "type": typ,
                            "text": src, "zh": zh}]}

    def test_truncated_ending_is_flagged(self):
        from llmtr import proofread

        src = ("This is a complete sentence about institutions and their evolution. " * 8
               + "The argument ends with a strong claim about state formation.")
        zh = "这是关于制度及其演变的完整陈述，" * 17 + "这一论证最后提出一个关于国家形成的强烈主张，并且"
        reasons = [i["reason"] for i in proofread.check(self._book(src, zh))]
        self.assertIn("疑似截断", reasons)

    def test_far_too_short_translation_is_flagged(self):
        from llmtr import proofread

        src = "A sentence of reasonable length about politics and history. " * 12
        zh = "这是一段明显不完整的译文，遗漏了原文的大部分内容。" * 5
        reasons = [i["reason"] for i in proofread.check(self._book(src, zh))]
        self.assertIn("译文偏短", reasons)

    def test_coverage_reports_pages_without_translation(self):
        from llmtr import proofread

        blocks = [
            {"id": "b0001", "page": "11", "type": "p", "text": "one", "zh": "一"},
            {"id": "b0002", "page": "12", "type": "p", "text": "two", "zh": ""},
            {"id": "b0003", "page": "12", "type": "figure", "text": "", "zh": ""},
        ]
        cov = proofread.coverage(blocks)
        self.assertEqual(cov["pages"], 2)
        self.assertEqual(cov["empty"], ["12"])

    def test_decade_written_in_chinese_numerals_is_kept(self):
        from llmtr import proofread

        src = ("The two regions shared formative experiences in the 1950s and 1960s, "
               "and that comparison still shapes the puzzles investigated here today.")
        zh = ("两个区域在二十世纪五十年代和六十年代共享形成性经历，"
              "这一比较至今仍塑造着我们此处考察的种种谜题。")
        reasons = [i["reason"] for i in proofread.check(self._book(src, zh))]
        self.assertNotIn("数字不一致", reasons)

    def test_dropped_years_are_still_flagged(self):
        from llmtr import proofread

        src = ("In 1978 and again in 1986 the ruling coalition replaced its model "
               "of political survival, reshaping the country's trajectory.")
        zh = "在这两个关键节点，统治联盟先后更换了政治生存模式，重塑了国家的发展轨迹。"
        reasons = [i["reason"] for i in proofread.check(self._book(src, zh))]
        self.assertIn("数字不一致", reasons)

    def test_long_citation_note_needs_no_translation(self):
        from llmtr import translate

        note = ("8 Iversen (1999); Boix (2000); Gourevitch (2003); Rodrik (1997); "
                "Schwartz (1998); Birdsall et al. (1999); Kapstein (2000); Rudra (2002).")
        self.assertTrue(translate._citation_only(note))
        self.assertFalse(translate._citation_only(
            "Such accounts overlook how domestic coalitions mediate external pressure."))


class StrayEnglishFix(unittest.TestCase):
    def test_guard_requires_words_actually_removed(self):
        from llmtr import translate

        old = "这一 dramatic 的转变使 successive 政权纷纷效仿，参见 Smith（1980）。"
        good = "这一急剧的转变使后继政权纷纷效仿，参见 Smith（1980）。"
        self.assertTrue(translate._term_fix_ok(old, good))
        self.assertFalse(translate._term_fix_ok(old, old))
        self.assertFalse(translate._term_fix_ok(old, good.replace("1980", "1990")))

    def test_stray_english_is_translated_in_place(self):
        from llmtr import translate

        blocks = [{"id": "b0001", "type": "p", "page": "1", "text": "A dramatic shift.",
                   "zh": "这一 dramatic 的转变使 successive 政权纷纷效仿，成为此后数十年的重要参照。"}]

        class Fake:
            key = "k"

            def chat(self, system, prompt, stop=None):
                return "这一急剧的转变使后继政权纷纷效仿，成为此后数十年的重要参照。"

        n = translate.fix_english_terms(Fake(), blocks, workers=1)
        self.assertEqual(n, 1)
        self.assertEqual(blocks[0]["zh"], "这一急剧的转变使后继政权纷纷效仿，成为此后数十年的重要参照。")


class PolishStageDefaults(unittest.TestCase):
    def test_polish_stage_is_wired(self):
        self.assertIn("润色", pipeline.STAGES)
        cfg = pipeline.load_config(None)
        self.assertTrue(cfg["polish"]["enabled"])
        self.assertLessEqual(cfg["chunk_words"], 2000)


if __name__ == "__main__":
    unittest.main()
